import os
import shutil
from datetime import datetime
import openpyxl

# form_path ='C:\\ARGOS RPA\\form\\금액 현황_한전.xlsx'
form_path ='C:\\Users\\vivans\\Desktop\\AI허브\\금액 현황_한전.xlsx'
# result_path = 'C:\\ARGOS RPA\\result\\'
result_path = 'C:\\Users\\vivans\\Desktop\\AI허브\\'

today = datetime.today()
today_month=today.strftime('%Y%m')



fname_path = result_path + '금액 현황_한전_' + today_month + '.xlsx'

fname = '금액 현황_한전_' + today_month + '.xlsx'

def main(customernum,elepay,elefund,elevat):
  flist = os.listdir(result_path)

  if fname in flist:
      wb = openpyxl.load_workbook(fname_path)
      ws = wb.active
  else:
      wb = openpyxl.load_workbook(form_path)
      ws = wb.active

  i = max((a.row for a in ws['A'] if a.value is not None))
  for i in range(5,i-1):
      if ws['C'+str(i)].value==customernum:
            a = elepay.replace(',', '')
            b = elefund.replace(',', '')
            c = elevat.replace(',', '')
            ws["E"+str(i)].value = int(a)
            ws["F"+str(i)].value = int(b)
            ws["H"+str(i)].value = int(c)
            x = int(a)+int(b)+int(c)
            ws["I"+str(i)].value = int(x)
            wb.save(fname_path)
            wb.close()
if __name__ == "__main__":
    main()
