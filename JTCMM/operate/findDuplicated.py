from openpyxl import load_workbook
import pandas as pd
import os


def do_pandas(file_path):
    path_ = file_path.replace(file_path.split("\\")[-1], "")
    file_name = file_path.split('\\')[-1].split('.')[0] + ".xlsx"
    df = pd.read_excel(file_path)
    not_error_df = df.drop_duplicates(keep=False)
    duplicated_error = df[df.duplicated(subset=['기관명'])]
    not_error_df.to_excel(path_ + file_name, index=False)
    if duplicated_error >= 1:
        duplicated_error.to_excel(path_ + file_path.split("\\")[-1].split(".")[0] + "_중복리스트" + ".xlsx", index=False)
    return path_ + file_name


def do_calc(no_error_file):
    wb = load_workbook(no_error_file)
    ws = wb.active
    for i in range(len(ws['A'])-1):
        student_QTY = ws['R' + str(i + 2)].value
        teacher_QTY = ws['S' + str(i + 2)].value
        if teacher_QTY > 0:
            new_teacher_QTY = teacher_QTY + 1
        else:
            new_teacher_QTY = teacher_QTY + 0
        ws['R' + str(i + 2)].value = student_QTY
        ws['S' + str(i + 2)].value = new_teacher_QTY
        ws['T' + str(i + 2)].value = new_teacher_QTY + student_QTY
    wb.save(no_error_file)
    return no_error_file


def main(file_path):
    do_calc(do_pandas(file_path))
    os.remove(file_path)


if __name__ == '__main__':
    main()
