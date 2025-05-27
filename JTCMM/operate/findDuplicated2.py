from openpyxl import load_workbook, Workbook
import pandas as pd
import os
import matplotlib.pyplot as plt



def do_pandas(file_path):
    path_ = file_path.replace(file_path.split("\\")[-1], "")
    file_name = file_path.split('\\')[-1].split('.')[0] + ".xlsx"
    df = pd.read_excel(file_path)
    not_error_df = df.drop_duplicates(keep=False)
    not_error_df.to_excel(path_ + file_name, index=False)
    error_df = df[df.duplicated(['주소'], keep=False) & df.duplicated(['상세주소'], keep=False)]
    error_df.to_excel(path_ + "temp.xlsx", index=False)
    return path_ + "temp.xlsx"


def delete_not_duplicated(file_path):
    wb = load_workbook(file_path)
    ws = wb.active
    counter = 0
    result = []
    title = []
    search_name = []
    for l in range(19):
        title.append(ws[chr(65+l)+"1"].value)
    title.append("중복 갯수")
    result.append(title)
    for i in range(2, ws.i):
        tmp = []
        tmp_1 = ws['C'+str(i)].value
        for j in range(i+1, ws.i + 1):
            tmp_2 = ws['C'+str(j)].value
            if tmp_1 == tmp_2:
                counter += 1
        if counter > 0:
            if tmp_1 not in search_name:
                for k in range(19):
                    tmp.append(ws[chr(65+k)+str(i)].value)
                tmp.append(counter+1)
                result.append(tmp)
                search_name.append(tmp_1)
                counter = 0
    os.remove(file_path)
    return result


def make_file(filepath, dupli_list):
    if len(dupli_list) == 1:
        return 0
    save_path = filepath.split(".")
    wb = Workbook()
    ws = wb.active
    for i in dupli_list:
        ws.append(i)
    wb.save(save_path[0] + "_중복리스트.xlsx")


def main(file_path):
    origin = do_pandas(file_path)
    duplicated_list = delete_not_duplicated(origin)
    make_file(file_path, duplicated_list)
    os.remove(file_path)


if __name__ == '__main__':
    main(r'C:\work\JTCMM\operate\230306 개인단말기.xls')
