from openpyxl import load_workbook


def main(file_path):
    wb = load_workbook(file_path)
    ws = wb.active
    all_QTY = 0
    for i in range(len(ws['A'])-1):
        all_QTY += ws['T'+str(i+2)].value
    return all_QTY


if __name__ == '__main__':
    main()
