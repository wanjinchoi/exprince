import openpyxl
import win32com.client as win32
import time
from openpyxl import load_workbook
from datetime import datetime
import os
from dateutil.relativedelta import relativedelta

NOW = datetime.now()
MONTH_NOW = datetime.now() - relativedelta(months=1)
#MONTH_DIRECTORY = r"C:\RPA\★월별출금내역\\" + str(MONTH_NOW.year) + "년"
MONTH_DIRECTORY = "C:\\work\\JTCMM\\rjfo\\효성cms\\" + str(MONTH_NOW.year) + "년"



def xls2xlsx(xls_path):
    if os.path.isfile(xls_path + "x"):
        os.remove(xls_path + "x")

    wb = win32.gencache.EnsureDispatch('Excel.Application').Workbooks.Open(xls_path)

    wb.SaveAs(xls_path + "x", FileFormat=51)
    wb.Close()

    win32.gencache.EnsureDispatch('Excel.Application').Application.Quit()
    time.sleep(3)

    xlsx_path = xls_path + "x"

    return xlsx_path



def extract_success_to_cms(path):
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    result_list = []
    for i in range(2, ws.i+1):
        tmp = []
        if ws.cell(i,10).value is None:
            pass
        else:
            if ("결제완료" in ws.cell(i,10).value)or ("결제실패" in ws.cell(i,10).value):
                for j in range(1,ws.max_column+1):
                        tmp.append(ws.cell(i,j).value)
            write_success_list(tmp)


def find_file_name(directory, date):
    listdir = os.listdir(directory)
    for i in listdir:
        if str(date) in i and "폐원" not in i and "직접" not in i:
            return i


def write_success_list(success_list):
    month = NOW - relativedelta(months=1)
    file_name = find_file_name(MONTH_DIRECTORY, str(month.month) + "월")
    wb = load_workbook(MONTH_DIRECTORY + "\\" + file_name)
    ws = wb.active
    failure_flag = False
    if len(success_list)>0:
        if success_list[9] == '재결제성공':
            ws.append(success_list)
        elif success_list[9] == '결제성공':
            ws.append(success_list)
        elif '결제실패' in success_list[9]:
            ws.append(success_list)
        elif success_list[9] == '결제완료':
            ws.append(success_list)
        wb.save(MONTH_DIRECTORY + "\\" + file_name)
    else:
        pass


def main(path):
    path = xls2xlsx(path)
    extract_success_to_cms(path)


if __name__ == '__main__':
    main(r'C:\work\JTCMM\rjfo\효성cms\exceldata.xls')