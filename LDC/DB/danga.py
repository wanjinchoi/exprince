import openpyxl


def main(change_xlsx):

    wb = openpyxl.load_workbook(change_xlsx)
    ws = wb.active
    #B열의 최대로우 구하기
    b_i = max((a.row for a in ws['B'] if a.value is not None))
    danga =[]
    for i in range(27, b_i+1):
        if ws['A'+str(i)].value is not None:
            danga.append(ws['F'+str(i)].value)

    list_str = list(map(str, danga))

    x = ','.join(list_str)

    return x







if __name__ == "__main__":
    main('C:\\Users\\vivans\\Desktop\\제출전\\VC231108003.xlsx')

