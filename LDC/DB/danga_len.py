import openpyxl
import re

def main(change_xlsx):

    wb = openpyxl.load_workbook(change_xlsx)
    ws = wb.active
    #B열의 최대로우 구하기
    b_i = max((a.row for a in ws['B'] if a.value is not None))
    danga =[]
    for i in range(27, b_i+1):
        if ws['A'+str(i)].value is not None:
            danga.append(ws['F'+str(i)].value)

    list_str = list(map(str, danga))
    #단가 길이 구하기
    le= len(list_str)
    ##################DISCOUNT 구하기
    #d열 최대로우 구하기
    d_i = max((a.row for a in ws['D'] if a.value is not None))
    discount =[]
    for j in range(27, d_i+1):
        if ws["D"+str(j)].value is None:
            pass
        else:
            if 'DISCOUNT' in str(ws["D"+str(j)].value):
                a = ws["D"+str(j)].value
                x = re.findall(r'\d+', a)
                x = ''.join(x)
                break
            else:
                x = 'nothing'

    final = str(le)+','+str(x)
    print(final)








if __name__ == "__main__":
    main('C:\\Users\\vivans\\Desktop\\제출전\\ZN231227006.xlsx')

