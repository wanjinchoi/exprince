import openpyxl




def main(change_xlsx):
    #엑셀읽기
    wb = openpyxl.load_workbook(change_xlsx)
    ws = wb.active
    remark =[]
    ad = []
    #===========================Qute Details 정보가져오기 ================================
    for i in range(2,ws.i+1):
        re = ws["A" + str(i)].value
        remark.append(re)

    for j in range(len(remark)):
        if remark[j] == None or remark[j]==' ':
            pass
        else:
            if 'REMARK' in remark[j] or'ITEM' in remark[j]:
                x = 'yes'
                break
            else:
                x ='nothing'

    if x =='yes':
        for i in range(2,ws.i+1):
            re = ws["A" + str(i)].value
            if re ==' ' or re == None:
                pass
            else:
                if "REMARK" in re or 'ITEM' in re:
                    if len(re) > 11:
                        re = re.replace('\n','')
                        x = re
                        break
                    else:
                        for j in range(i, ws.i + 1):
                            re = ws["A" + str(j)].value
                            ad.append(re)
        q = [s for s in ad if s !=' ' and s != None]
        w = ''
        for index, item in enumerate(q):
            if index == 1:
                if (item != None) and ('REMARK' not in item) and (item != ' '):
                    w += item + '\n'
            else:
                if (item != None) and ('REMARK' not in item) and (item != ' '):
                    w += item
        if not q:
           return x
        else:
           x = "REMARK: " + '\n' + w
           return x
    else:
        return x
    wb.close()







if __name__ == "__main__":
    main('C:\\Users\\vivans\\Desktop\\제출전\\UO231206002.xlsx')

