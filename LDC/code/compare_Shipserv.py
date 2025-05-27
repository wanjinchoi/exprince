import os

import openpyxl
from datetime import datetime

name_path='C:\\work\\LDC\\Name.xlsx'
today = datetime.today()

#ARGOS
# load_path = 'C:\\work\\LDC\\form\\'+today.strftime('%m%d')+'V_Name.xlsx'
# nulack_filepath = 'C:\\work\\LDC\\nulack\\'+today.strftime('%m%d')+'nulack_Name.xlsx'
# nulack_path = 'C:\\work\\LDC\\nulack\\'
# form_path = 'C:\\work\\LDC\\form\\nulack.xlsx'

#LDC
load_path = 'C:\\ArgosRPA\\form\\'+today.strftime('%m%d')+'Name.xlsx'
nulack_filepath = 'C:\\ArgosRPA\\nulack\\'+today.strftime('%m%d')+'nulack_Name.xlsx'
nulack_path = 'C:\\ArgosRPA\\nulack\\'
form_path = 'C:\\ArgosRPA\\form\\nulack.xlsx'


def main(paper_num):

    if paper_num =='' or paper_num is None:
        b = 'no'
        return b

    else:
        wb_ck = openpyxl.load_workbook(load_path)
        ws_ck = wb_ck.active
        complete_list = [cell.value for cell in ws_ck['A']]
        if paper_num in complete_list:
            b = 'yes'
            return b
        else:
            nulack_exist = os.listdir(nulack_path)
            if today.strftime('%m%d')+'nulack_Name.xlsx' in nulack_exist:
                wb = openpyxl.load_workbook(nulack_filepath)
                ws = wb.active
                max_row = max((a.row for a in ws['A'] if a.value is not None))
                if ws['A'+str(max_row)].value == paper_num:
                    pass
                else:
                    ws['A' + str(max_row + 1)].value = paper_num
                    wb.save(nulack_filepath)
                wb.close()
            else:
                wb = openpyxl.load_workbook(form_path)
                ws = wb.active
                max_row = max((a.row for a in ws['A'] if a.value is not None))
                ws['A'+str(max_row+1)].value = paper_num
                wb.save(nulack_filepath)
                wb.close()
            b='no'
            return b
if __name__ == "__main__":
    main()

