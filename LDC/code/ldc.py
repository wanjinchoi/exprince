import re
import glob
import openpyxl
from datetime import datetime
import os
#가공하기 위한 양식파일
#template_path = 'C:\\work\\LDC\\form.xlsx'
##가공한 엑셀파일 경로
#result_path = 'C:\\work\\LDC\\result.xlsx'
#file_path = 'C:\\work\\LDC\\예시excel파일\\'


today = datetime.today()
fileday = today.strftime('%Y%m%d')
template_path = 'C:\\ArgosRPA\\form\\form.xlsx'
result_path ='C:\\ArgosRPA\\result\\result.xlsx'
file_path = 'C:\\ArgosRPA\\'+fileday+'\\'


def main(comment,subject):
    flist = sorted(glob.glob(file_path + '*.xlsx'), key=os.path.getmtime)
    r_len = len(flist) - 1
    file = flist[r_len]
    #다운한 오리지널 엑셀파일 활성화
    wb_or = openpyxl.load_workbook(file)
    ws_or = wb_or.active
    #양식 엑세파일 활성화
    wb_fm = openpyxl.load_workbook(template_path)
    ws_fm = wb_fm.active

    if subject == 'nothing':
        pass
    else:
        ws_fm['B2'].value = '+'
        ws_fm['D2'].value = subject
        wb_fm.save(result_path)
    if comment == 'nothing':
        pass
    elif comment != 'nothing':
        if len(comment) > 250:
            a = len(comment)
            b = comment
            c = a / 250
            z = []
            z.append(b[0: 250])
            for i in range(1, int(c) + 1):
                y = b[(i * 250) + 1:(i + 1) * 250 + 1]
                z.append(y)
            j = 3
            for d3 in z:
                ws_fm["D" + str(j)].value = d3
                ws_fm["B" + str(j)].value = "-"
                wb_fm.save(result_path)
                j += 1
        else:
            ws_fm['B3'].value = '-'
            x = comment
            z = x.replace('  ', '')
            c = z.replace('\n', ' ')
            ws_fm['D3'].value = c
            wb_fm.save(result_path)


    for i in range(2,ws_or.max_row+1):
        d_max_row = max((a.row for a in ws_fm['D'] if a.value is not None))
        #I
        if ws_or["O"+str(1)].value == 'Eq. Section Name':
            o = ws_or["O" + str(i)].value
        else:
            o = ws_or["I" + str(i)].value
        #J
        if ws_or["P" + str(1)].value == 'Eq. Section Description':
            p = ws_or["P" + str(i)].value
        else:
            p = ws_or["J" + str(i)].value
        #K
        if ws_or["Q" + str(1)].value == 'Eq. Section Manufacturer':
            q = ws_or["Q" + str(i)].value
        else:
            q = ws_or["K" + str(i)].value
        #L
        if ws_or["R" + str(1)].value == 'Eq. Section Model Number':
            r = ws_or["R" + str(i)].value
        else:
            r = ws_or["L" + str(i)].value
        #M
        if ws_or["S" + str(1)].value == 'Eq. Section Rating':
             s = ws_or["S" + str(i)].value
        else:
            s = ws_or["M" + str(i)].value
        #N
        if ws_or["T" + str(1)].value == 'Eq. Section Serial Number':
            t = ws_or["T" + str(i)].value
        else:
            t = ws_or["N" + str(i)].value
        #O
        if ws_or["U" + str(1)].value == 'Eq. Section Drawing Number':
            u = ws_or["U" + str(i)].value
        else:
            u = ws_or["O" + str(i)].value
        #P
        if ws_or["V" + str(1)].value == 'Eq. Section Department Type':
            v = ws_or["V" + str(i)].value
        else:
            v = ws_or["P" + str(i)].value
        if o is None:
            o = 'Section Name: ,@'
        else:
            o = 'Section Name: ' + o+ ',@'
        if p is None:
            p = 'Description: ,@'
        else:
            p = 'Description: ' + p + ',@'
        if q is None:
            q = 'Manufacturer: ,@'
        else:
            q = 'Manufacturer: ' + q+ ',@'
        if r is None:
            r = 'Model Number: ,@'
        else:
            r = 'Model Number: ' + r + ',@'
        if s is None:
            s = 'Rating: ,@'
        else:
            s = 'Rating: ' + s + ',@'
        if t is None:
            t = 'Serial Number: ,@'
        else:
            t = 'Serial Number: ' + t + ',@'
        if u is None:
            u = 'Drawing Number: ,@'
        else:
            u = 'Drawing Number: ' + u + ',@'
        if v is None:
            v = 'Department Type: ,@'
        else:
            v = 'Department Type: ' + v + ',@'
        eq_section = o + p + q + r + s + t + u + v
        eq_section = eq_section.replace('Section Name: ,@', '').replace('Description: ,@', '').replace('Manufacturer: ,@','').replace('Model Number: ,@', '').replace('Rating: ,@', '').replace('Serial Number: ,@', '').replace('Drawing Number: ,@','').replace('Department Type: ,@', '')
        eq_section = eq_section.replace('@', ' ').replace('Section ','')
        eq_section = eq_section.rstrip(", ")

        if ws_or["O"+str(1)].value == 'Eq. Section Name':
            o = ws_or["O" + str(i+1)].value
        else:
            o = ws_or["I" + str(i+1)].value
        #J
        if ws_or["P" + str(1)].value == 'Eq. Section Description':
            p = ws_or["P" + str(i+1)].value
        else:
            p = ws_or["J" + str(i+1)].value
        #K
        if ws_or["Q" + str(1)].value == 'Eq. Section Manufacturer':
            q = ws_or["Q" + str(i+1)].value
        else:
            q = ws_or["K" + str(i+1)].value
        #L
        if ws_or["R" + str(1)].value == 'Eq. Section Model Number':
            r = ws_or["R" + str(i+1)].value
        else:
            r = ws_or["L" + str(i+1)].value
        #M
        if ws_or["S" + str(1)].value == 'Eq. Section Rating':
             s = ws_or["S" + str(i+1)].value
        else:
            s = ws_or["M" + str(i+1)].value
        #N
        if ws_or["T" + str(1)].value == 'Eq. Section Serial Number':
            t = ws_or["T" + str(i+1)].value
        else:
            t = ws_or["N" + str(i+1)].value
        #O
        if ws_or["U" + str(1)].value == 'Eq. Section Drawing Number':
            u = ws_or["U" + str(i+1)].value
        else:
            u = ws_or["O" + str(i+1)].value
        #P
        if ws_or["V" + str(1)].value == 'Eq. Section Department Type':
            v = ws_or["V" + str(i+1)].value
        else:
            v = ws_or["P" + str(i+1)].value

        if o is None:
            o = 'Section Name: ,@'
        else:
            o = 'Section Name: ' + o + ',@'
        if p is None:
            p = 'Description: ,@'
        else:
            p = 'Description: ' +p + ',@'
        if q is None:
            q = 'Manufacturer: ,@'
        else:
            q = 'Manufacturer: ' + q+ ',@'
        if r is None:
            r = 'Model Number: ,@'
        else:
            r = 'Model Number: ' + r + ',@'
        if s is None:
            s = 'Rating: ,@'
        else:
            s = 'Rating: ' + s + ',@'
        if t is None:
            t = 'Serial Number: ,@'
        else:
            t = 'Serial Number: ' + t + ',@'
        if u is None:
            u = 'Drawing Number: ,@'
        else:
            u = 'Drawing Number: ' + u + ',@'
        if v is None:
            v = 'Department Type: ,@'
        else:
            v = 'Department Type: ' + v + ',@'
        eq_section2 = o + p + q + r + s + t + u + v
        eq_section2 = eq_section2.replace('Section Name: ,@', '').replace('Description: ,@', '').replace('Manufacturer: ,@', '').replace('Model Number: ,@', '').replace('Rating: ,@', '').replace('Serial Number: ,@', '').replace('Drawing Number: ,@', '').replace('Department Type: ,@', '')
        eq_section2 = eq_section2.replace('@', ' ').replace('Section ','')
        eq_section2 = eq_section2.rstrip(", ")

        if i==2 and eq_section2 == '':
            d_max_row = max((a.row for a in ws_fm['D'] if a.value is not None))
            if len(eq_section)> 240:
                    eq_len = len(eq_section)
                    eq_sh = eq_len / 240
                    z = []
                    z.append(eq_section[0: 240])
                    x = 1
                    for m in range(1, int(eq_sh) + 1):
                        yz = eq_section[(m * 240) + 1:(m + 1) * 240 + 1]
                        z.append(yz)
                    for dz in z:
                        if dz == '':
                            break
                        d_max_row = max((a.row for a in ws_fm['D'] if a.value is not None))
                        if x == 1:
                            ws_fm["B" + str(d_max_row + 1)].value = '-'
                            ws_fm["D" + str(d_max_row + 1)].value = dz
                            wb_fm.save(result_path)
                            x += 1
                        else:
                            ws_fm["B" + str(d_max_row + 1)].value = '-'
                            ws_fm["D" + str(d_max_row + 1)].value = dz
                            wb_fm.save(result_path)
                            x += 1
            else:
                ws_fm["B" + str(d_max_row + 1)].value = '-'
                ws_fm["D" + str(d_max_row + 1)].value = eq_section
                wb_fm.save(result_path)

            d_max_row = max((a.row for a in ws_fm['D'] if a.value is not None))
            ws_fm["A" + str(d_max_row + 1)].value = ws_or['A' + str(i)].value
            ws_fm["C" + str(d_max_row + 1)].value = ws_or["C" + str(i)].value
            if ws_or["E" + str(1)].value == 'Item Description':
                b = ws_or["E" + str(i)].value
            else:
                b = b = ws_or["D" + str(i)].value
            if b is not None:
                reb = re.compile(r'^\[[a-zA-Z]+\]')
                a = reb.split(b)
                if len(a) > 1:
                    b = a[1]
                else:
                    b = a[0]
                b = b.replace('  ', '')
                b = b.replace('\n', ' ')
                b = b.replace('\0', '')
                b = b.replace('\r', '')
                b = b.replace('\t', ' ')
                if len(b) > 250:
                    b_len = len(b)
                    b_sh = b_len / 250
                    za = []
                    za.append(b[0: 250])
                    for m in range(1, int(b_sh) + 1):
                        yz = b[(m * 250) + 1:(m + 1) * 250 + 1]
                        za.append(yz)
                        x = 1
                    for dz in za:
                        if dz == '':
                            break
                        d_max_row = max((a.row for a in ws_fm['D'] if a.value is not None))

                        if x == 1:
                            #QTY
                            if ws_or["H1"].value =="Qty":
                                qty = ws_or["H" + str(i)].value
                            else:
                                qty = ws_or["F" + str(i)].value
                            ws_fm["E" + str(d_max_row + 1)].value = qty
                            #Unit
                            ws_fm["F" + str(d_max_row + 1)].value = ws_or["G" + str(i)].value
                            ws_fm["D" + str(d_max_row + 1)].value = dz
                            wb_fm.save(result_path)
                            x += 1
                        else:
                            ws_fm["B" + str(d_max_row + 1)].value = '='
                            ws_fm["D" + str(d_max_row + 1)].value = dz
                            wb_fm.save(result_path)
                            x += 1
                else:
                    ws_fm["D" + str(d_max_row + 1)].value = b
                    if ws_or["H1"].value == "Qty":
                        qty = ws_or["H" + str(i)].value
                    else:
                        qty = ws_or["F" + str(i)].value
                    ws_fm["E" + str(d_max_row + 1)].value = qty
                    ws_fm["F" + str(d_max_row + 1)].value = ws_or["G" + str(i)].value
            d_max_row = max((a.row for a in ws_fm['D'] if a.value is not None))
            if ws_or["J" + str(1)].value == 'Buyer Comments':
                c = ws_or["J" + str(i)].value
            else:
                c = ws_or["E" + str(i)].value
            if c is not None:
                ah = c.replace('\n', ' ')
                ah2 = ah.replace('\t', ' ')
                if len(ah) > 250:
                    ah_len = len(ah)
                    bc = ah
                    cd = ah_len / 250
                    za = []
                    za.append(bc[0: 250])
                    for i in range(1, int(cd) + 1):
                        yz = bc[(i * 250) + 1:(i + 1) * 250 + 1]
                        za.append(yz)
                    for dz in za:
                        if dz == '':
                            break
                        d_max_row = max((a.row for a in ws_fm['D'] if a.value is not None))
                        ws_fm["B" + str(d_max_row + 1)].value = '='
                        ws_fm["D" + str(d_max_row + 1)].value = dz
                        wb_fm.save(result_path)
                else:
                    ws_fm["B" + str(d_max_row + 1)].value = '='
                    ws_fm["D" + str(d_max_row + 1)].value = ah2
            wb_fm.save(result_path)
        else:
            if eq_section == eq_section2:
                if i ==2:
                    d_max_row = max((a.row for a in ws_fm['D'] if a.value is not None))
                    if len(eq_section) > 240:
                        eq_len = len(eq_section)
                        eq_sh = eq_len / 240
                        z = []
                        z.append(eq_section[0:240])
                        x = 1
                        for m in range(1, int(eq_sh) + 1):
                            yz = eq_section[(m * 240) + 1:(m + 1) * 240 + 1]
                            z.append(yz)
                        for dz in z:
                            if dz == '':
                                break
                            d_max_row = max((a.row for a in ws_fm['D'] if a.value is not None))
                            if x == 1:
                                ws_fm["B" + str(d_max_row + 1)].value = '-'
                                ws_fm["D" + str(d_max_row + 1)].value = dz
                                wb_fm.save(result_path)
                                x += 1
                            else:
                                ws_fm["B" + str(d_max_row + 1)].value = '-'
                                ws_fm["D" + str(d_max_row + 1)].value = dz
                                wb_fm.save(result_path)
                                x += 1
                    else:
                        ws_fm["B" + str(d_max_row + 1)].value = '-'
                        ws_fm["D" + str(d_max_row + 1)].value = eq_section
                        wb_fm.save(result_path)
                    d_max_row = max((a.row for a in ws_fm['D'] if a.value is not None))
                    ws_fm["A" + str(d_max_row + 1)].value = ws_or['A' + str(i)].value
                    ws_fm["C" + str(d_max_row + 1)].value = ws_or["C" + str(i)].value
                    if ws_or["E" + str(1)].value == 'Item Description':
                        b = ws_or["E" + str(i)].value
                    else:
                        b = b = ws_or["D" + str(i)].value
                    if b is not None:
                        reb = re.compile(r'^\[[a-zA-Z]+\]')
                        a = reb.split(b)
                        if len(a) > 1:
                            b = a[1]
                        else:
                            b = a[0]
                        b = b.replace('  ', '')
                        b = b.replace('\n', ' ')
                        b = b.replace('\0', '')
                        b = b.replace('\r', '')
                        b = b.replace('\t', ' ')
                        if len(b)>250:
                            b_len = len(b)
                            b_sh = b_len/ 250
                            za = []
                            za.append(b[0: 250])
                            for m in range(1, int(b_sh) + 1):
                                yz = b[(m * 250) + 1:(m + 1) * 250 + 1]
                                za.append(yz)
                                x = 1
                            for dz in za:
                                if dz == '':
                                    break
                                d_max_row = max((a.row for a in ws_fm['D'] if a.value is not None))

                                if x == 1:
                                    if ws_or["H1"].value == "Qty":
                                        qty = ws_or["H" + str(i)].value
                                    else:
                                        qty = ws_or["F" + str(i)].value
                                    ws_fm["E" + str(d_max_row + 1)].value = qty
                                    ws_fm["F" + str(d_max_row + 1)].value = ws_or["G" + str(i)].value
                                    ws_fm["D" + str(d_max_row + 1)].value = dz
                                    wb_fm.save(result_path)
                                    x += 1
                                else:
                                    ws_fm["B" + str(d_max_row + 1)].value = '='
                                    ws_fm["D" + str(d_max_row + 1)].value = dz
                                    wb_fm.save(result_path)
                                    x += 1

                        else:
                            ws_fm["D" + str(d_max_row + 1)].value = b
                            if ws_or["H1"].value =="Qty":
                                qty = ws_or["H" + str(i)].value
                            else:
                                qty = ws_or["F" + str(i)].value
                            ws_fm["E" + str(d_max_row + 1)].value = qty
                            ws_fm["F" + str(d_max_row + 1)].value = ws_or["G" + str(i)].value
                    wb_fm.save(result_path)
                    d_max_row = max((a.row for a in ws_fm['D'] if a.value is not None))
                    if ws_or["J"+str(1)].value =='Buyer Comments':
                        c = ws_or["J" + str(i)].value
                    else:
                        c = ws_or["E"+str(i)].value
                    if c is not None:
                        ah = c.replace('\n', ' ')
                        ah2 = ah.replace('\t', ' ')
                        if len(ah) > 250:
                            ah_len = len(ah)
                            bc = ah
                            cd = ah_len / 250
                            za = []
                            za.append(bc[0: 250])
                            for i in range(1, int(cd) + 1):
                                yz = bc[(i * 250) + 1:(i + 1) * 250 + 1]
                                za.append(yz)
                            for dz in za:
                                if dz == '':
                                    break
                                d_max_row = max((a.row for a in ws_fm['D'] if a.value is not None))
                                ws_fm["B" + str(d_max_row + 1)].value = '='
                                ws_fm["D" + str(d_max_row + 1)].value = dz
                                wb_fm.save(result_path)
                        else:
                            ws_fm["B" + str(d_max_row + 1)].value = '='
                            ws_fm["D" + str(d_max_row + 1)].value = ah2
                    wb_fm.save(result_path)
                else:
                    d_max_row = max((a.row for a in ws_fm['D'] if a.value is not None))
                    ws_fm["A" + str(d_max_row + 1)].value = ws_or['A' + str(i)].value
                    ws_fm["C" + str(d_max_row + 1)].value = ws_or["C" + str(i)].value
                    if ws_or["E" + str(1)].value == 'Item Description':
                        b = ws_or["E" + str(i)].value
                    else:
                        b = b = ws_or["D" + str(i)].value
                    if b is not None:
                        reb = re.compile(r'^\[[a-zA-Z]+\]')
                        a = reb.split(b)
                        if len(a) > 1:
                            b = a[1]
                        else:
                            b = a[0]
                        b = b.replace('  ', '')
                        b = b.replace('\n', ' ')
                        b = b.replace('\0', '')
                        b = b.replace('\r', '')
                        b = b.replace('\t', ' ')
                        if len(b)>250:
                            b_len = len(b)
                            b_sh = b_len/ 250
                            za = []
                            za.append(b[0: 250])
                            for m in range(1, int(b_sh) + 1):
                                yz = b[(m * 250) + 1:(m + 1) * 250 + 1]
                                za.append(yz)
                                x = 1
                            for dz in za:
                                if dz == '':
                                    break
                                d_max_row = max((a.row for a in ws_fm['D'] if a.value is not None))

                                if x == 1:
                                    if ws_or["H1"].value == "Qty":
                                        qty = ws_or["H" + str(i)].value
                                    else:
                                        qty = ws_or["F" + str(i)].value
                                    ws_fm["E" + str(d_max_row + 1)].value = qty
                                    ws_fm["F" + str(d_max_row + 1)].value = ws_or["G" + str(i)].value
                                    ws_fm["D" + str(d_max_row + 1)].value = dz
                                    wb_fm.save(result_path)
                                    x += 1
                                else:
                                    ws_fm["B" + str(d_max_row + 1)].value = '='
                                    ws_fm["D" + str(d_max_row + 1)].value = dz
                                    wb_fm.save(result_path)
                                    x += 1

                        else:
                            ws_fm["D" + str(d_max_row + 1)].value = b
                            if ws_or["H1"].value =="Qty":
                                qty = ws_or["H" + str(i)].value
                            else:
                                qty = ws_or["F" + str(i)].value
                            ws_fm["E" + str(d_max_row + 1)].value = qty
                            ws_fm["F" + str(d_max_row + 1)].value = ws_or["G" + str(i)].value
                    d_max_row = max((a.row for a in ws_fm['D'] if a.value is not None))
                    if ws_or["J"+str(1)].value =='Buyer Comments':
                        c = ws_or["J" + str(i)].value
                    else:
                        c = ws_or["E"+str(i)].value
                    if c is not None:
                        ah = c.replace('\n', ' ')
                        ah2 = ah.replace('\t', ' ')
                        if len(ah) > 250:
                            ah_len = len(ah)
                            bc = ah
                            cd = ah_len / 250
                            za = []
                            za.append(bc[0: 250])
                            for i in range(1, int(cd) + 1):
                                yz = bc[(i * 250) + 1:(i + 1) * 250 + 1]
                                za.append(yz)
                            for dz in za:
                                if dz == '':
                                    break
                                d_max_row = max((a.row for a in ws_fm['D'] if a.value is not None))
                                ws_fm["B" + str(d_max_row + 1)].value = '='
                                ws_fm["D" + str(d_max_row + 1)].value = dz
                                wb_fm.save(result_path)
                        else:
                            ws_fm["B" + str(d_max_row + 1)].value = '='
                            ws_fm["D" + str(d_max_row + 1)].value = ah2
                    wb_fm.save(result_path)
            else:
                d_max_row = max((a.row for a in ws_fm['D'] if a.value is not None))
                if eq_section2 == '':
                    d_max_row = max((a.row for a in ws_fm['D'] if a.value is not None))
                    ws_fm["A" + str(d_max_row + 1)].value = ws_or['A' + str(i)].value
                    ws_fm["C" + str(d_max_row + 1)].value = ws_or["C" + str(i)].value
                    if ws_or["E" + str(1)].value == 'Item Description':
                        b = ws_or["E" + str(i)].value
                    else:
                        b = b = ws_or["D" + str(i)].value
                    if b is not None:
                        reb = re.compile(r'^\[[a-zA-Z]+\]')
                        a = reb.split(b)
                        if len(a) > 1:
                            b = a[1]
                        else:
                            b = a[0]
                        b = b.replace('  ', '')
                        b = b.replace('\n', ' ')
                        b = b.replace('\0', '')
                        b = b.replace('\r', '')
                        b = b.replace('\t', ' ')
                        if len(b) > 250:
                            b_len = len(b)
                            b_sh = b_len / 250
                            za = []
                            za.append(b[0: 250])
                            for m in range(1, int(b_sh) + 1):
                                yz = b[(m * 250) + 1:(m + 1) * 250 + 1]
                                za.append(yz)
                                x = 1
                            for dz in za:
                                if dz == '':
                                    break
                                d_max_row = max((a.row for a in ws_fm['D'] if a.value is not None))

                                if x == 1:
                                    if ws_or["H1"].value == "Qty":
                                        qty = ws_or["H" + str(i)].value
                                    else:
                                        qty = ws_or["F" + str(i)].value
                                    ws_fm["E" + str(d_max_row + 1)].value = qty
                                    ws_fm["F" + str(d_max_row + 1)].value = ws_or["G" + str(i)].value
                                    ws_fm["D" + str(d_max_row + 1)].value = dz
                                    wb_fm.save(result_path)
                                    x += 1
                                else:
                                    ws_fm["B" + str(d_max_row + 1)].value = '='
                                    ws_fm["D" + str(d_max_row + 1)].value = dz
                                    wb_fm.save(result_path)
                                    x += 1

                        else:
                            ws_fm["D" + str(d_max_row + 1)].value = b
                            if ws_or["H1"].value =="Qty":
                                qty = ws_or["H" + str(i)].value
                            else:
                                qty = ws_or["F" + str(i)].value
                            ws_fm["E" + str(d_max_row + 1)].value = qty
                            ws_fm["F" + str(d_max_row + 1)].value = ws_or["G" + str(i)].value
                    d_max_row = max((a.row for a in ws_fm['D'] if a.value is not None))
                    if ws_or["J"+str(1)].value =='Buyer Comments':
                        c = ws_or["J" + str(i)].value
                    else:
                        c = ws_or["E"+str(i)].value
                    if c is not None:
                        ah = c.replace('\n', ' ')
                        ah2 = ah.replace('\t', ' ')
                        if len(ah) > 250:
                            ah_len = len(ah)
                            bc = ah
                            cd = ah_len / 250
                            za = []
                            za.append(bc[0: 250])
                            for i in range(1, int(cd) + 1):
                                yz = bc[(i * 250) + 1:(i + 1) * 250 + 1]
                                za.append(yz)
                            for dz in za:
                                if dz == '':
                                    break
                                d_max_row = max((a.row for a in ws_fm['D'] if a.value is not None))
                                ws_fm["B" + str(d_max_row + 1)].value = '='
                                ws_fm["D" + str(d_max_row + 1)].value = dz
                                wb_fm.save(result_path)
                        else:
                            ws_fm["B" + str(d_max_row + 1)].value = '='
                            ws_fm["D" + str(d_max_row + 1)].value = ah2
                    wb_fm.save(result_path)
                    d_max_row = max((a.row for a in ws_fm['D'] if a.value is not None))
                else:
                    if ws_fm["D" + str(d_max_row)].value == eq_section:
                        pass
                    else:
                        if len(eq_section) > 240:
                            eq_len = len(eq_section)
                            eq_sh = eq_len / 240
                            z = []
                            z.append(eq_section[0: 240])
                            x = 1
                            for m in range(1, int(eq_sh) + 1):
                                yz = eq_section[(m * 240) + 1:(m + 1) * 240 + 1]
                                z.append(yz)
                            for dz in z:
                                if dz == '':
                                    break
                                d_max_row = max((a.row for a in ws_fm['D'] if a.value is not None))
                                if x == 1:
                                    ws_fm["B" + str(d_max_row + 1)].value = '-'
                                    ws_fm["D" + str(d_max_row + 1)].value = dz
                                    wb_fm.save(result_path)
                                    x += 1
                                else:
                                    ws_fm["B" + str(d_max_row + 1)].value = '-'
                                    ws_fm["D" + str(d_max_row + 1)].value = dz
                                    wb_fm.save(result_path)
                                    x += 1
                        else:
                            ws_fm["B" + str(d_max_row + 1)].value = '-'
                            ws_fm["D" + str(d_max_row + 1)].value = eq_section
                            wb_fm.save(result_path)
                    d_max_row = max((a.row for a in ws_fm['D'] if a.value is not None))
                    ws_fm["A" + str(d_max_row + 1)].value = ws_or['A' + str(i)].value
                    ws_fm["C" + str(d_max_row + 1)].value = ws_or["C" + str(i)].value
                    if ws_or["E" + str(1)].value == 'Item Description':
                        b = ws_or["E" + str(i)].value
                    else:
                        b = b = ws_or["D" + str(i)].value
                    if b is not None:
                        reb = re.compile(r'^\[[a-zA-Z]+\]')
                        a = reb.split(b)
                        if len(a) > 1:
                            b = a[1]
                        else:
                            b = a[0]
                        b = b.replace('  ', '')
                        b = b.replace('\n', ' ')
                        b = b.replace('\0', '')
                        b = b.replace('\r', '')
                        b = b.replace('\t', ' ')
                        if len(b) > 250:
                            b_len = len(b)
                            b_sh = b_len / 250
                            za = []
                            za.append(b[0: 250])
                            for m in range(1, int(b_sh) + 1):
                                yz = b[(m * 250) + 1:(m + 1) * 250 + 1]
                                za.append(yz)
                                x = 1
                            for dz in za:
                                if dz == '':
                                    break
                                d_max_row = max((a.row for a in ws_fm['D'] if a.value is not None))

                                if x == 1:
                                    if ws_or["H1"].value == "Qty":
                                        qty = ws_or["H" + str(i)].value
                                    else:
                                        qty = ws_or["F" + str(i)].value
                                    ws_fm["E" + str(d_max_row + 1)].value = qty
                                    ws_fm["F" + str(d_max_row + 1)].value = ws_or["G" + str(i)].value
                                    ws_fm["D" + str(d_max_row + 1)].value = dz
                                    wb_fm.save(result_path)
                                    x += 1
                                else:
                                    ws_fm["B" + str(d_max_row + 1)].value = '='
                                    ws_fm["D" + str(d_max_row + 1)].value = dz
                                    wb_fm.save(result_path)
                                    x += 1

                        else:
                            ws_fm["D" + str(d_max_row + 1)].value = b
                            if ws_or["H1"].value =="Qty":
                                qty = ws_or["H" + str(i)].value
                            else:
                                qty = ws_or["F" + str(i)].value
                            ws_fm["E" + str(d_max_row + 1)].value = qty
                            ws_fm["F" + str(d_max_row + 1)].value = ws_or["G" + str(i)].value
                    d_max_row = max((a.row for a in ws_fm['D'] if a.value is not None))
                    if ws_or["J"+str(1)].value =='Buyer Comments':
                        c = ws_or["J" + str(i)].value
                    else:
                        c = ws_or["E"+str(i)].value
                    if c is not None:
                        ah = c.replace('\n', ' ')
                        ah2 = ah.replace('\t', ' ')
                        if len(ah) > 250:
                            ah_len = len(ah)
                            bc = ah
                            cd = ah_len / 250
                            za = []
                            za.append(bc[0: 250])
                            for i in range(1, int(cd) + 1):
                                yz = bc[(i * 250) + 1:(i + 1) * 250 + 1]
                                za.append(yz)
                            for dz in za:
                                if dz == '':
                                    break
                                d_max_row = max((a.row for a in ws_fm['D'] if a.value is not None))
                                ws_fm["B" + str(d_max_row + 1)].value = '='
                                ws_fm["D" + str(d_max_row + 1)].value = dz
                                wb_fm.save(result_path)
                        else:
                            ws_fm["B" + str(d_max_row + 1)].value = '='
                            ws_fm["D" + str(d_max_row + 1)].value = ah2
                    wb_fm.save(result_path)
                    d_max_row = max((a.row for a in ws_fm['D'] if a.value is not None))
                    if eq_section2 == '':
                        pass
                    else:
                        ws_fm["B" + str(d_max_row+1)].value = '-'
                        ws_fm["D" + str(d_max_row+1)].value = eq_section2
                    wb_fm.save(result_path)

if __name__ == "__main__":
    main()
