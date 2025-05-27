import openpyxl
from datetime import datetime
from dateutil.relativedelta import relativedelta
import os

today = datetime.today()
oneday_ago = today -relativedelta(days=1)
twoday_ago = today - relativedelta(days=2)



#form_path = 'C:\\work\\LDC\\form\\'
#load_path = 'C:\\work\\LDC\\form\\' + today.strftime('%m%d') + 'Name.xlsx'
#oneday_ago_path = 'C:\\work\\LDC\\form\\'+oneday_ago.strftime('%m%d')+'Name.xlsx'
#twoday_ago_path = 'C:\\work\\LDC\\form\\' + twoday_ago.strftime('%m%d') + 'Name.xlsx'


#Nozzle
form_path = 'C:\\ArgosRPA\\Nozzle\\form\\'
load_path = 'C:\\ArgosRPA\\Nozzle\\form\\' + today.strftime('%m%d') + 'NOName.xlsx'
oneday_ago_path = 'C:\\ArgosRPA\\Nozzle\\form\\'+oneday_ago.strftime('%m%d')+'NOName.xlsx'
twoday_ago_path = 'C:\\ArgosRPA\\Nozzle\\form\\' + twoday_ago.strftime('%m%d') + 'NOName.xlsx'

#@meps
# form_path='C:\\ArgosRPA\\form\\Name.xlsx'
# load_path ='C:\\ArgosRPA\\meps\\form\\'+today.strftime('%m%d')+'m_name.xlsx'
# oneday_ago_path = 'C:\\ArgosRPA\\meps\\form\\'+oneday_ago.strftime('%m%d')+'m_name.xlsx'
# twoday_ago_path ='C:\\ArgosRPA\\meps\\form\\'+twoday_ago.strftime('%m%d')+'m_name.xlsx'


#@kinntienk
# form_path='C:\\ArgosRPA\\form\\Name.xlsx'
# load_path ='C:\\ArgosRPA\\kinnetik\\form\\'+today.strftime('%m%d')+'k_Name.xlsx'
# oneday_ago_path ='C:\\ArgosRPA\\kinnetik\\form\\'+oneday_ago.strftime('%m%d')+'k_Name.xlsx'
# twoday_ago_path ='C:\\ArgosRPA\\kinnetik\\form\\'+twoday_ago.strftime('%m%d')+'k_Name.xlsx'

#@Vgroup
# form_path='C:\\ArgosRPA\\vgroup\\form\\V_Name.xlsx'
# load_path = 'C:\\ArgosRPA\\vgroup\\form\\'+today.strftime('%m%d')+'V_Name.xlsx'
# twoday_ago_path='C:\\ArgosRPA\\vgroup\\form\\'+twoday_ago.strftime('%m%d')+'V_Name.xlsx'

#@Prosureship
# form_path='C:\\ArgosRPA\\form\\Name.xlsx'
# load_path ='C:\\ArgosRPA\\porcure\\form\\'+today.strftime('%m%d')+'p_name.xlsx'
# oneday_ago_path ='C:\\ArgosRPA\\porcure\\form\\'+oneday_ago.strftime('%m%d')+'p_name.xlsx'
# twoday_ago_path ='C:\\ArgosRPA\\porcure\\form\\'+twoday_ago.strftime('%m%d')+'p_name.xlsx'

checklist=r'{checklist}'
if os.path.isfile(load_path) == True:
    wb_ck = openpyxl.load_workbook(load_path)
    ws_ck = wb_ck.active
else:
    wb_ck = openpyxl.load_workbook(oneday_ago_path)
    ws_ck = wb_ck.active
    wb_ck.save(load_path)
#if os.path.isfile(twoday_ago_path) == True:
    #os.remove(twoday_ago_path)
a =[]
b= []
for i in range(2, ws_ck.max_row+2):
    complete = ws_ck["A"+str(i)].value
    a.append(complete)
    doit = ws_ck["B"+str(i)].value
    b.append(doit)

if checklist in a:
    result = checklist
    print(result)
else:
    if checklist not in b:
        ws_ck["B2"].value = checklist
        wb_ck.save(load_path)
        result ="doit"
        print(result)
a = dict()
a["a"] = b

