import os

import openpyxl
import pandas as pd
#name_path='C:\\ArgosRPA\\form\\Name.xlsx'
excel_path= 'C:\\work\\LDC\\meps\\meps_csv\\'

def main(checklist):
    try:
        df = pd.read_csv(excel_path+'meps.csv', encoding='cp949',sep='\t')
        df.to_excel(excel_path+'meps.xlsx', index=False)
        wb = openpyxl.load_workbook(excel_path + 'meps.xlsx')
        ws = wb.active
        for i in range(1, ws.i + 1):
            if '혻' in ws["A" + str(i)].value:
                ws["A" + str(i)].value = ws["A" + str(i)].value.replace('혻', ' ')
        wb.save(excel_path + 'meps.xlsx')
    except:
        wb = openpyxl.load_workbook(excel_path+'meps.xlsx')
        ws = wb.active
        wb.remove(ws)
        wb.create_sheet('Sheet1')
        wb.save(excel_path+'meps.xlsx')




if __name__ == "__main__":
    main('aa')

