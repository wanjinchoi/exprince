import openpyxl
import pandas as pd
import os
import glob

# form_path='C:\\work\\LDC\\form.xlsx'
# excel_path='C:\\work\\LDC\\excel\\xlsx_file\\'
# result_path='C:\\work\\LDC\\v_result.xlsx'
# except_path ='C:\\work\\LDC\\except.xlsx'





form_path='C:\\ArgosRPA\\form\\form.xlsx'
excel_path='C:\\ArgosRPA\\vgroup\\excel\\xlsx_file\\'
result_path='C:\\ArgosRPA\\vgroup\\result\\v_result.xlsx'
except_path ='C:\\ArgosRPA\\vgroup\\form\\except.xlsx'








def main(subject,vcom,excelpath):
    csv_file = os.listdir(excelpath)
    for file in csv_file:
        try:
            if file =='itemno.csv':
                df = pd.read_csv(excelpath + '\\' + file,encoding='cp949',sep='\t')
                df.to_excel(excel_path+'itemno.xlsx', index=False)
            if file =='ldcpartname.csv':
                df = pd.read_csv(excelpath + '\\' + file,encoding='UTF8',sep='\t')
                df.to_excel(excel_path+'ldcpartname.xlsx', index=False)
            if file == 'ldcpartno.csv':
                df = pd.read_csv(excelpath + '\\' + file,encoding='cp949',sep='\t')
                df.to_excel(excel_path+'ldcpartno.xlsx', index=False)
            if file =='ldcqty.csv':
                df = pd.read_csv(excelpath + '\\' + file,encoding='cp949',sep='\t')
                df.to_excel(excel_path+'ldcqty.xlsx', index=False)
            if file =='ldcunit.csv':
                df = pd.read_csv(excelpath + '\\' + file,encoding='cp949',sep='\t')
                df.to_excel(excel_path+'ldcunit.xlsx', index=False)
            if file =='ldcclib.csv':
                df = pd.read_csv(excelpath + '\\' + 'ldcclib.csv', encoding='UTF8', sep='\t')
                df.to_excel(excel_path + 'ldcclib.xlsx', index=False)
            if file =='ldcdrp.csv':
                df = pd.read_csv(excelpath + '\\' + file, encoding='UTF8', sep='\t')
                df.to_excel(excel_path + 'ldcdrp.xlsx', index=False)


        except:
            file2= file.replace('.csv','.xlsx')
            wb = openpyxl.load_workbook(except_path)
            ws = wb.active
            wb.remove(ws)
            wb.create_sheet('Sheet1')
            wb.save(excel_path + file2)

    wb_fm = openpyxl.load_workbook(form_path)
    ws_fm = wb_fm.active
    ws_fm["B2"].value ='+'
    ws_fm["D2"].value = subject
    ws_fm["B3"].value = '-'
    if len(vcom) > 250:
         a = len(vcom)
         b = vcom
         c = a/250
         z=[]
         z.append(b[0: 250])
         for i in range(1, int(c)+1):
             y = b[(i*250)+1:(i+1)*250+1]
             z.append(y)
         j = 3
         for d3 in z:
             if 'Serial: Component' in vcom:
                 d3 = d3.replace('Serial:','')
                 if 'Maker: Seria' in vcom:
                     vcom = vcom.replace('Maker:', '')
             else:
                 pass
             ws_fm["D"+str(j)].value = d3
             ws_fm["B"+str(j)].value ="-"
             wb_fm.save(result_path)
             j+=1
    else:
        if 'Serial: Component' in vcom:
            vcom = vcom.replace('Serial:', '')
            if 'Maker: Seria' in vcom:
                vcom = vcom.replace('Maker:', '')
        else:
            pass
        ws_fm['D3'].value = vcom
        wb_fm.save(result_path)

    #[os.remove(f) for f in glob.glob(excelpath+'*.csv')]
    wb_fm = openpyxl.load_workbook(result_path)
    ws_fm = wb_fm.active
    wb_pa = openpyxl.load_workbook(excel_path+'ldcpartname.xlsx')
    ws_pa = wb_pa.active
    #clib(excel_path)
    wb_clib = openpyxl.load_workbook(excel_path + 'ldcclib.xlsx')
    ws_clib = wb_clib.active
    wb_drp = openpyxl.load_workbook(excel_path+'ldcdrp.xlsx')
    ws_drp = wb_drp.active
    if (ws_clib['A1'].value is None)and(ws_drp['A1'].value is None):
        for i in range(1, ws_pa.max_row + 1):
            d_i = max((a.row for a in ws_fm['D'] if a.value is not None))
            wb_it = openpyxl.load_workbook(excel_path + 'itemno.xlsx')
            ws_it = wb_it.active
            wb_pno = openpyxl.load_workbook(excel_path + 'ldcpartno.xlsx')
            ws_pno = wb_pno.active
            wb_qty = openpyxl.load_workbook(excel_path + 'ldcqty.xlsx')
            ws_qty = wb_qty.active
            wb_unit = openpyxl.load_workbook(excel_path + 'ldcunit.xlsx')
            ws_unit = wb_unit.active
            # itemno입력
            ws_fm["A" + str(d_i + 1)].value = ws_it["A" + str(i)].value
            wb_fm.save(result_path)
            # partno입력
            if 'nothing'in str(ws_pno["A" + str(i)].value):
                pass
            elif ws_pno["A" + str(i)].value =='na':
                pass
            else:
                ws_fm["C" + str(d_i + 1)].value = ws_pno["A" + str(i)].value
            # 수량입력
            ws_fm["E" + str(d_i + 1)].value = ws_qty['A' + str(i)].value
            # 단위입력
            ws_fm['F' + str(d_i + 1)].value = ws_unit['A' + str(i)].value
            wb_fm.save(result_path)
            b = ws_pa['A' + str(i)].value
            if b is not None:
                ax = b.replace('\n', ' ')
                ax2 = ax.replace('\0', '')
                av = ax2.replace('\r', '')
                axc = av.replace('\t', '')
                ws_fm['D' + str(d_i + 1)].value = axc
                wb_fm.save(result_path)
                d_i = max((a.row for a in ws_fm['D'] if a.value is not None))

            else:
                break
    else:
        j=1
        for i in range(1, ws_pa.max_row + 1):
            d_i = max((a.row for a in ws_fm['D'] if a.value is not None))
            wb_it = openpyxl.load_workbook(excel_path + 'itemno.xlsx')
            ws_it = wb_it.active
            wb_pno = openpyxl.load_workbook(excel_path + 'ldcpartno.xlsx')
            ws_pno = wb_pno.active
            wb_qty = openpyxl.load_workbook(excel_path + 'ldcqty.xlsx')
            ws_qty = wb_qty.active
            wb_unit = openpyxl.load_workbook(excel_path + 'ldcunit.xlsx')
            ws_unit = wb_unit.active
            # itemno입력
            ws_fm["A" + str(d_i + 1)].value = ws_it["A" + str(i)].value
            # partno입력
            if 'nothing' in str(ws_pno["A" + str(i)].value):
                pass
            elif ws_pno["A" + str(i)].value =='na':
                pass
            else:
                if "째" in str(ws_pno["A" + str(i)].value):
                    ws_pno["A" + str(i)].value = ws_pno["A" + str(i)].value.replace("째",'°')
                    ws_fm["C" + str(d_i + 1)].value = ws_pno["A" + str(i)].value
                else:
                    ws_fm["C" + str(d_i + 1)].value = ws_pno["A" + str(i)].value
            # 수량입력
            ws_fm["E" + str(d_i + 1)].value = ws_qty['A' + str(i)].value
            # 단위입력
            ws_fm['F' + str(d_i + 1)].value = ws_unit['A' + str(i)].value
            wb_fm.save(result_path)
            b = str(ws_pa['A' + str(i)].value)
            if b is not None:
                ax = b.replace('\n', ' ')
                ax2 = ax.replace('\0', '')
                av = ax2.replace('\r', '')
                axc = av.replace('\t', '')
                ws_fm['D'+str(d_i+1)].value=axc
                wb_fm.save(result_path)
                d_i = max((a.row for a in ws_fm['D'] if a.value is not None))
                if ('nothing'in str(ws_clib["A"+str(i)].value))or(str(ws_clib["A"+str(i)].value) is None):
                    if (ws_drp['A' + str(i)].value != 'nothing') and (ws_drp['A' + str(i)].value is not None) and (ws_drp['A' + str(i)].value != ',nothing'):
                        ws_fm['B' + str(d_i + 1)].value = '='
                        ws_fm["D" + str(d_i + 1)].value = ws_drp['A' + str(i)].value
                        wb_fm.save(result_path)
                elif str(ws_clib['A'+str(i)].value):
                        pass
                else:
                    clib = ws_clib["A" + str(i)].value
                    if 'C:\\ARGOSRPA\\VGROUP\\EXCEL\\LDCCLIB.CSV' in str(clib):
                        clib= clib.replace('C:\\ARGOSRPA\\VGROUP\\EXCEL\\LDCCLIB.CSV', '')
                    if len(clib) > 250:
                        a = len(clib)
                        b = clib
                        c = a / 250
                        z = []
                        z.append(b[0: 250])
                        for i in range(1, int(c) + 1):
                            y = b[(i * 250) + 1:(i + 1) * 250 + 1]
                            z.append(y)
                        for rclib in z:
                            d_i = max((a.row for a in ws_fm['D'] if a.value is not None))
                            ws_fm['B' + str(d_i + 1)].value = '='
                            ws_fm['D' + str(d_i + 1)].value = rclib
                            wb_fm.save(result_path)
                    else:
                        ws_fm['B' + str(d_i + 1)].value = '='
                        ws_fm['D' + str(d_i + 1)].value = ws_clib["A" + str(i)].value
                        wb_fm.save(result_path)
                    if ('nothing' in str(ws_drp['A' + str(i)].value)) or (str(ws_drp['A' + str(i)].value) is None) :
                        pass
                    else:
                        if ws_fm["D" + str(d_i+1)].value is None:
                            ws_fm["D" + str(d_i + 1)].value = str(ws_drp['A' + str(i)].value)
                            wb_fm.save(result_path)
                        else:
                            ws_fm["D" + str(d_i+1)].value =ws_fm["D" + str(d_i+1)].value+', '+str(ws_drp['A' + str(i)].value)
                            clib2 = ws_fm["D" + str(d_i+1)].value
                            if 'C:\\ARGOSRPA\\VGROUP\\EXCEL\\LDCCLIB.CSV' in clib2:
                                clib2 = clib2.replace('C:\\ARGOSRPA\\VGROUP\\EXCEL\\LDCCLIB.CSV', '')
                            if len(clib2) > 250:
                                a = len(clib2)
                                b = clib2
                                c = a / 250
                                z = []
                                z.append(b[0: 250])
                                for i in range(1, int(c) + 1):
                                    y = b[(i * 250) + 1:(i + 1) * 250 + 1]
                                    z.append(y)
                                d_i = max((a.row for a in ws_fm['D'] if a.value is not None))
                                for rclib2 in z:
                                    ws_fm['D'+str(d_i)].value =rclib2
                                    d_i +=1
                                    wb_fm.save(result_path)
                            wb_fm.save(result_path)

            else:
                break

    ## 엑셀파일 열어서
    wb_re = openpyxl.load_workbook(result_path)
    ws_re = wb_re.active

    # 모든 셀 순회하면서 값 치환
    for row in ws_re.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and "°C" in cell.value:
                cell.value = cell.value.replace("°C", "'C")

    # 저장
    wb_re.save(result_path)
    wb_re.close()

if __name__ == "__main__":
    main('aa','bb',r'C:\work\LDC\excel')


