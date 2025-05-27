import openpyxl
#가공하기 위한 양식파일
#template_path = 'C:\\work\\LDC\\form.xlsx'
##가공한 엑셀파일 경로
#result_path = 'C:\\work\\LDC\\result.xlsx'


template_path = 'C:\\ArgosRPA\\form\\form.xlsx'
result_path ='C:\\ArgosRPA\\result\\result.xlsx'




def main(comment,subject,file):
    #다운한 오리지널 엑셀파일 활성화
    wb_or = openpyxl.load_workbook(file)
    ws_or = wb_or.active
    #양식 엑세파일 활성화
    wb_fm = openpyxl.load_workbook(template_path)
    ws_fm = wb_fm.active


    #### O열과 P열에서 Name이 있기때문에 O열인지 P열인지 첫번째 대분류 O열인경우
    if ws_or['O1'].value =='Eq. Section Name':
        if (ws_or['O2'].value is None) and (ws_or['O3'].value is None):
            wb_fm = openpyxl.load_workbook(template_path)
            ws_fm = wb_fm.active
            if subject =='nothing':
                ws_fm['B2'].value = '+'
            else:
                ws_fm['B2'].value = '+'
                ws_fm['D2'].value = subject
            if comment == 'nothing':
                    pass
            elif comment !='nothing':
                if len(comment) > 250:
                    a = len(comment)
                    b = comment
                    c = a / 250
                    z = []
                    z.append(b[0: 250])
                    for i in range(1, int(c) + 1):
                        y = b[(i * 250) + 1:(i + 1) * 250 + 1]
                        z.append(y)
                    j = 3
                    for d3 in z:
                        ws_fm["D" + str(j)].value = d3
                        ws_fm["B" + str(j)].value = "-"
                        wb_fm.save(result_path)
                        j += 1
                else:
                    ws_fm['B3'].value ='-'
                    ws_fm['D3'].value =comment
                    wb_fm.save(result_path)
            d_i = max((a.row for a in ws_fm['D'] if a.value is not None))
            for i in range(2, 3):
                if ws_or["Q" + str(i)].value is None:
                    pass
                else:
                    ws_fm['D'+ str(d_i+1)].value = 'ManuFacturer:' + ws_or['Q' + str(i)].value
                    wb_fm.save(result_path)
                if ws_or["R" + str(i)].value is None:
                    pass
                else:
                    if ws_fm['D'+str(d_i+1)].value is not None:
                        ws_fm['D'+str(d_i+1)].value = ws_fm['D'+str(d_i+1)].value + ',' + 'ModelNumber:' + ws_or["R" + str(i)].value
                        wb_fm.save(result_path)
                    else:
                        ws_fm['D'+str(d_i+1)].value ='ModelNumber:' + ws_or["R" + str(i)].value
                        wb_fm.save(result_path)
                if ws_or["V" + str(i)].value is None:
                    pass
                else:
                   if ws_fm['D'+str(d_i+1)].value is not None:
                        ws_fm['D'+str(d_i+1)].value = ws_fm['D'+str(d_i+1)].value + ',' + 'DepartmentType:' + ws_or["V" + str(i)].value
                        wb_fm.save(result_path)
                   else:
                       ws_fm['D'+str(d_i+1)].value = 'DepartmentType:' + ws_or["V" + str(i)].value
                       wb_fm.save(result_path)
                if ws_or["T" + str(i)].value is None:
                    pass
                else:
                    if ws_fm['D'+str(d_i+1)].value is not None:
                        ws_fm['D'+str(d_i+1)].value = ws_fm['D'+str(d_i+1)].value + ',' + 'SerialNumber:' + ws_or["T" + str(i)].value
                    else:
                        ws_fm['D'+str(d_i+1)].value = 'SerialNumber:' + ws_or["T" + str(i)].value
                        wb_fm.save(result_path)

            if ws_fm['D'+str(d_i+1)].value is not None:
                ws_fm['B'+str(d_i+1)].value ='-'
            else:
                pass
            wb_fm.save(result_path)
            e = ws_fm.i + 1
            i=2
            for j in range(2, ws_or.i+1):
                e = ws_fm.i + 1
                ws_fm["A"+str(e)].value = ws_or["A"+str(i)].value
                if ws_or["C" + str(i)].value == 'N/A':
                    pass
                else:
                    ws_fm["C"+str(e)].value = ws_or["C"+str(i)].value
                b = ws_or["E" + str(i)].value
                if b is not None :
                    ax = b.replace('\n', ' ')
                    ax2 = ax.replace('\0', '')
                    av = ax2.replace('\r', '')
                    axc = av.replace('\t', '')
                    ws_fm["D"+str(e)].value = axc
                ws_fm["D"+str(e)].value = ws_or["E" + str(i)].value
                ws_fm["E" + str(e)].value = ws_or["F" + str(i)].value
                ws_fm["F" + str(e)].value = ws_or["G" + str(i)].value
                c = ws_or["J" + str(i)].value
                if c is not None:
                    ah = c.replace('\n',' ')
                    ah2 = ah.replace('\t',' ')
                    if len(ah)> 250:
                        ah_len = len(ah)
                        bc = ah
                        cd = ah_len / 250
                        za = []
                        za.append(bc[0: 250])
                        for i in range(1, int(cd) + 1):
                            yz = bc[(i * 250) + 1:(i + 1) * 250 + 1]
                            za.append(yz)
                        for dz in za:
                            if dz == '':
                                break
                            d_i = max((a.row for a in ws_fm['D'] if a.value is not None))
                            ws_fm["B" +str(d_i+1)].value = '='
                            ws_fm["D" +str(d_i+1)].value = dz
                            wb_fm.save(result_path)
                    else:
                        ws_fm["B" + str(e + 1)].value = '='
                        ws_fm["D" + str(e + 1)].value = ah2
                wb_fm.save(result_path)
                i +=1



        else:
            wb_or = openpyxl.load_workbook(file)
            ws_or = wb_or.active
            wb_fm = openpyxl.load_workbook(template_path)
            ws_fm = wb_fm.active
            if subject =='nothing':
                ws_fm['B2'].value = '+'
            else:
                ws_fm['B2'].value = '+'
                ws_fm['D2'].value = subject
            if comment == 'nothing':
                pass
            elif comment !='nothing':
                if len(comment) > 250:
                    a = len(comment)
                    b = comment
                    c = a / 250
                    z = []
                    z.append(b[0: 250])
                    for i in range(1, int(c) + 1):
                        y = b[(i * 250) + 1:(i + 1) * 250 + 1]
                        z.append(y)
                    j = 3
                    for d3 in z:
                        ws_fm["D" + str(j)].value = d3
                        ws_fm["B" + str(j)].value = "-"
                        wb_fm.save(result_path)
                        j += 1
                else:
                    ws_fm['B3'].value ='-'
                    ws_fm['D3'].value =comment
                    wb_fm.save(result_path)
            d_i = max((a.row for a in ws_fm['D'] if a.value is not None))
            for i in range(2, 3):
                if ws_or["Q"+str(i)].value is None:
                    pass
                else:
                    ws_fm['D'+ str(d_i+1)].value = 'ManuFacturer:' + ws_or['Q' + str(i)].value
                    wb_fm.save(result_path)
                if ws_or["R" +str(i)].value is None:
                    pass
                else:
                    if ws_fm['D'+str(d_i+1)].value is not None:
                        ws_fm['D'+str(d_i+1)].value = ws_fm['D'+str(d_i+1)].value + ',' + 'ModelNumber:' + ws_or["R" + str(i)].value
                        wb_fm.save(result_path)
                    else:
                        ws_fm['D'+str(d_i+1)].value ='ModelNumber:' + ws_or["R" + str(i)].value
                        wb_fm.save(result_path)
                if ws_or["V" + str(i)].value is None:
                    pass
                else:
                   if ws_fm['D'+str(d_i+1)].value is not None:
                        ws_fm['D'+str(d_i+1)].value = ws_fm['D'+str(d_i+1)].value + ',' + 'DepartmentType:' + ws_or["V" + str(i)].value
                        wb_fm.save(result_path)
                   else:
                       ws_fm['D'+str(d_i+1)].value = 'DepartmentType:' + ws_or["V" + str(i)].value
                       wb_fm.save(result_path)
                if ws_or["T" + str(i)].value is None:
                    pass
                else:
                    if ws_fm['D'+str(d_i+1)].value is not None:
                        ws_fm['D'+str(d_i+1)].value = ws_fm['D'+str(d_i+1)].value + ',' + 'SerialNumber:' + ws_or["T" + str(i)].value
                    else:
                        ws_fm['D'+str(d_i+1)].value = 'SerialNumber:' + ws_or["T" + str(i)].value
                        wb_fm.save(result_path)


            if ws_fm['D'+str(d_i+1)].value is not None:
                ws_fm['B'+str(d_i+1)].value ='-'
            else:
                pass
            wb_fm.save(result_path)
            e=ws_fm.i+1
            i=2
            for j in range(2, ws_or.i + 1):
                if (ws_or['O' + str(i)].value == ws_or['O' + str(i + 1)].value) and (ws_or['O' + str(i)].value is not None):
                    #print('-이름이 같음')
                    e = ws_fm.i + 1
                    # 앞선 방법과 동일하게 처리
                    ws_fm["B" + str(e)].value = "-"
                    ws_fm["D" + str(e)].value = ws_or["O" + str(i)].value
                    wb_fm.save(result_path)
                    e = ws_fm.i + 1
                    for j in range(2, ws_or.i+1):
                        e = ws_fm.i + 1
                        ws_fm["A"+str(e)].value = ws_or["A"+str(i)].value
                        if ws_or["C"+str(i)].value =='N/A':
                            pass
                        else:
                            ws_fm["C"+str(e)].value = ws_or["C"+str(i)].value
                        if len(ws_or["E"+str(i)].value) >250:
                            a= ws_or["E"+str(i)].value
                            under = a[0:250]
                            under2 = under.replace('\n','')
                            under3 = under2.replace('\t','')
                            under4 = under3.replace('\0','')
                            under5 = under4.replace('\r','')
                            over = a[250:]
                            over2 = over.replace('\n','')
                            over3 = over2.replace('\t','')
                            over4 = over3.replace('\0','')
                            over5 = over4.replace('\r','')
                            ws_fm["D"+str(e)].value = under5
                            ws_fm["B" + str(e+1)].value = "="
                            ws_fm["D"+str(e+1)].value =over5
                            ws_fm["E" + str(e)].value = ws_or["F" + str(i)].value
                            ws_fm["F" + str(e)].value = ws_or["G" + str(i)].value
                            c = ws_or["J" + str(i)].value
                            if c is not None:
                                ah = c.replace('\n', ' ')
                                ah2 = ah.replace('\t', ' ')
                                if len(ah) > 250:
                                    ah_len = len(ah)
                                    bc = ah
                                    cd = ah_len / 250
                                    za = []
                                    za.append(bc[0: 250])
                                    for i in range(1, int(cd) + 1):
                                        yz = bc[(i * 250) + 1:(i + 1) * 250 + 1]
                                        za.append(yz)
                                    for dz in za:
                                        if dz == '':
                                            break
                                        d_i = max((a.row for a in ws_fm['D'] if a.value is not None))
                                        ws_fm["B" + str(d_i + 1)].value = '='
                                        ws_fm["D" + str(d_i + 1)].value = dz
                                        wb_fm.save(result_path)
                                else:
                                     ws_fm["B" + str(e + 1)].value = '='
                                     ws_fm["D" + str(e + 1)].value = ah2
                            wb_fm.save(result_path)
                            i +=1
                        else:
                            b = ws_or["E" + str(i)].value
                            if b is not None:
                                ax = b.replace('\n', ' ')
                                ax2 = ax.replace('\0', '')
                                av = ax2.replace('\r', '')
                                axc = av.replace('\t', '')
                                ws_fm["D" + str(e)].value = axc
                            ws_fm["E" + str(e)].value = ws_or["F" + str(i)].value
                            ws_fm["F" + str(e)].value = ws_or["G" + str(i)].value
                            wb_fm.save(result_path)
                            c = ws_or["J" + str(i)].value
                            if c is not None:
                                ah = c.replace('\n', ' ')
                                ah2 = ah.replace('\t', ' ')
                                if len(ah) > 250:
                                    ah_len = len(ah)
                                    bc = ah
                                    cd = ah_len / 250
                                    za = []
                                    za.append(bc[0: 250])
                                    for i in range(1, int(cd) + 1):
                                        yz = bc[(i * 250) + 1:(i + 1) * 250 + 1]
                                        za.append(yz)
                                    for dz in za:
                                        if dz == '':
                                            break
                                        d_i = max((a.row for a in ws_fm['D'] if a.value is not None))
                                        ws_fm["B" + str(d_i + 1)].value = '='
                                        ws_fm["D" + str(d_i + 1)].value = dz
                                        wb_fm.save(result_path)
                                else:
                                    ws_fm["B" + str(e + 1)].value = '='
                                    ws_fm["D" + str(e + 1)].value = ah2
                            wb_fm.save(result_path)
                            i +=1
                        if ws_or["A" + str(i)].value is None:
                            break



                if (ws_or['O' + str(i)].value != ws_or['O' + str(i + 1)].value) and (ws_or['O' + str(i)].value is not None):
                    #print("- 이름이 다름")
                    e = ws_fm.i + 1
                    ws_fm["B"+str(e)].value ="-"
                    ws_fm["D"+str(e)].value =ws_or["O"+str(i)].value
                    wb_fm.save(result_path)
                    e = ws_fm.i+1
                    ws_fm["A"+str(e)].value = ws_or["A"+str(i)].value
                    if ws_or["C"+str(i)].value =='N/A':
                        pass
                    else:
                        ws_fm["C"+str(e)].value = ws_or["C"+str(i)].value
                    ws_fm["E"+str(e)].value = ws_or["F"+str(i)].value
                    ws_fm["F" + str(e)].value = ws_or["G" + str(i)].value
                    if len(ws_or["E"+str(i)].value) > 250:
                        a=ws_or["E"+str(i)].value
                        under = a[0:250]
                        under2 = under.replace('\n', ' ')
                        under3 = under2.replace('\t', '')
                        under4 = under3.replace('\0', '')
                        under5 = under4.replace('\r', '')
                        over = a[250:]
                        over2 = over.replace('\n', ' ')
                        over3 = over2.replace('\t', '')
                        over4 = over3.replace('\0', '')
                        over5 = over4.replace('\r', '')
                        ws_fm["D" + str(e)].value = under5
                        ws_fm["B"+ str(e+1)].value ="="
                        ws_fm["D"+str(e+1)].value =over5
                        c = ws_or["J" + str(i)].value
                        if c is not None:
                            ah = c.replace('\n', ' ')
                            ah2 = ah.replace('\t', ' ')
                            if len(ah) > 250:
                                ah_len = len(ah)
                                bc = ah
                                cd = ah_len / 250
                                za = []
                                za.append(bc[0: 250])
                                for i in range(1, int(cd) + 1):
                                    yz = bc[(i * 250) + 1:(i + 1) * 250 + 1]
                                    za.append(yz)
                                for dz in za:
                                    if dz == '':
                                        break
                                    d_i = max((a.row for a in ws_fm['D'] if a.value is not None))
                                    ws_fm["B" + str(d_i + 1)].value = '='
                                    ws_fm["D" + str(d_i + 1)].value = dz
                                    wb_fm.save(result_path)
                            else:
                                ws_fm["B" + str(e + 1)].value = '='
                                ws_fm["D" + str(e + 1)].value = ah2
                        wb_fm.save(result_path)
                        e=ws_fm.i+1
                        i +=1
                    b = ws_or["E" + str(i)].value
                    if b is not None :
                        ax = b.replace('\n', ' ')
                        ax2 = ax.replace('\0', '')
                        av = ax2.replace('\r', '')
                        axc = av.replace('\t', '')
                        ws_fm["D"+str(e)].value = axc
                    ws_fm["D"+str(e)].value = ws_or["E" + str(i)].value
                    ws_fm["E"+str(e)].value = ws_or["F"+str(i)].value
                    ws_fm["F" + str(e)].value = ws_or["G" + str(i)].value
                    c = ws_or["J" + str(i)].value
                    if c is not None:
                        ah = c.replace('\n', ' ')
                        ah2 = ah.replace('\t', ' ')
                        if len(ah) > 250:
                            ah_len = len(ah)
                            bc = ah
                            cd = ah_len / 250
                            za = []
                            za.append(bc[0: 250])
                            for i in range(1, int(cd) + 1):
                                yz = bc[(i * 250) + 1:(i + 1) * 250 + 1]
                                za.append(yz)
                            for dz in za:
                                if dz == '':
                                    break
                                d_i = max((a.row for a in ws_fm['D'] if a.value is not None))
                                ws_fm["B" + str(d_i + 1)].value = '='
                                ws_fm["D" + str(d_i + 1)].value = dz
                                wb_fm.save(result_path)
                        else:
                            ws_fm["B" + str(e + 1)].value = '='
                            ws_fm["D" + str(e + 1)].value = ah2
                    wb_fm.save(result_path)
                    i+=1




#########################################################################################################################
    #### O열과 P열에서 Name이 있기때문에 O열인지 P열인지 첫번째 대분류 P열인경우
    if ws_or['P1'].value =='Eq. Section Name':
        if (ws_or['P2'].value is None) and (ws_or['P3'].value is None):
            wb_fm = openpyxl.load_workbook(template_path)
            ws_fm = wb_fm.active
            if subject =='nothing':
                ws_fm['B2'].value = '+'
            else:
                ws_fm['B2'].value = '+'
                ws_fm['D2'].value = subject
            if comment == 'nothing':
                pass
            elif comment != 'nothing':
                if len(comment) > 250:
                    a = len(comment)
                    b = comment
                    c = a / 250
                    z = []
                    z.append(b[0: 250])
                    for i in range(1, int(c) + 1):
                        y = b[(i * 250) + 1:(i + 1) * 250 + 1]
                        z.append(y)
                    j = 3
                    for d3 in z:
                        ws_fm["D" + str(j)].value = d3
                        ws_fm["B" + str(j)].value = "-"
                        wb_fm.save(result_path)
                        j += 1
                else:
                    ws_fm['B3'].value = '-'
                    ws_fm['D3'].value = comment
            d_i = max((a.row for a in ws_fm['D'] if a.value is not None))
            wb_fm.save(result_path)
            for i in range(2, 3):
                if ws_or["R" + str(i)].value is None:
                    pass
                else:
                    ws_fm['D'+str(d_i+1)].value = 'ManuFacturer:' + ws_or['R' + str(i)].value
                    wb_fm.save(result_path)
                if ws_or["S" + str(i)].value is None:
                    pass
                else:
                    if ws_fm['D'+str(d_i+1)].value is None:
                        ws_fm['D' + str(d_i + 1)].value = ''
                        ws_fm['D'+str(d_i+2)].value = ws_fm['D'+str(d_i+1)].value + ',' + 'ModelNumber:' + ws_or["S" + str(i)].value
                        wb_fm.save(result_path)
                    else:
                        ws_fm['D'+str(d_i+1)].value = ws_fm['D'+str(d_i+1)].value + ',' + 'ModelNumber:' + ws_or["S" + str(i)].value
                        wb_fm.save(result_path)
                if ws_or["W" + str(i)].value is None:
                    pass
                else:
                    if ws_fm['D'+str(d_i+1)].value is None:
                        ws_fm['D'+str(d_i+1)].value = ''
                        ws_fm['D'+str(d_i+2)].value = ws_fm['D'+str(d_i+1)].value + ',' + 'DepartmentType:' + ws_or["W" + str(i)].value
                        wb_fm.save(result_path)
                    else:
                        ws_fm['D'+str(d_i+1)].value = ws_fm['D'+str(d_i+1)].value + ',' + 'DepartmentType:' + ws_or["W" + str(i)].value
                        wb_fm.save(result_path)
                if ws_or["U" + str(i)].value is None:
                    pass
                else:
                    if ws_fm['D'+str(d_i+1)].value is None:
                        ws_fm['D'+str(d_i+1)].value = ''
                        ws_fm['D'+str(d_i+2)].value = ws_fm['D'+str(d_i+1)].value + ',' + 'SerialNumber:' + ws_or["U" + str(i)].value
                        break
                    else:
                        ws_fm['D'+str(d_i+1)].value = ws_fm['D'+str(d_i+1)].value + ',' + 'SerialNumber:' + ws_or["U" + str(i)].value
                        wb_fm.save(result_path)
            if ws_fm['D'+str(d_i+1)].value is not None:
                ws_fm['B'+str(d_i+1)].value ='-'
            else:
                pass
            wb_fm.save(result_path)
            e = ws_fm.i + 1
            i = 2
            for j in range(2, ws_or.i + 1):
                e = ws_fm.i + 1
                ws_fm["A" + str(e)].value = ws_or["A" + str(i)].value
                ws_fm["C" + str(e)].value = ws_or["C" + str(i)].value
                ws_fm["E" + str(e)].value = ws_or["F" + str(i)].value
                ws_fm["F" + str(e)].value = ws_or["G" + str(i)].value
                b = ws_or["E" + str(i)].value
                if b is not None :
                    ax = b.replace('\n', ' ')
                    ax2 = ax.replace('\0', '')
                    av = ax2.replace('\r', '')
                    axc = av.replace('\t', '')
                    ws_fm["D"+str(e)].value = axc
                c = ws_or["J" + str(i)].value
                if c is not None:
                    ah = c.replace('\n',' ')
                    ah2 = ah.replace('\t',' ')
                    if len(ah)> 250:
                        ah_len = len(ah)
                        bc = ah
                        cd = ah_len / 250
                        za = []
                        za.append(bc[0: 250])
                        for i in range(1, int(cd) + 1):
                            yz = bc[(i * 250) + 1:(i + 1) * 250 + 1]
                            za.append(yz)
                        for dz in za:
                            if dz == '':
                                break
                            d_i = max((a.row for a in ws_fm['D'] if a.value is not None))
                            ws_fm["B" +str(d_i+1)].value = '='
                            ws_fm["D" +str(d_i+1)].value = dz
                            wb_fm.save(result_path)
                    else:
                        ws_fm["B" + str(e + 1)].value = '='
                        ws_fm["D" + str(e + 1)].value = ah2
                wb_fm.save(result_path)
                i += 1
        else:
            wb_or = openpyxl.load_workbook(file)
            ws_or = wb_or.active
            wb_fm = openpyxl.load_workbook(template_path)
            ws_fm = wb_fm.active
            if subject =='nothing':
                ws_fm['B2'].value = '+'
            else:
                ws_fm['B2'].value = '+'
                ws_fm['D2'].value = subject
            if comment == 'nothing':
                pass
            elif comment != 'nothing':
                if len(comment) > 250:
                    a = len(comment)
                    b = comment
                    c = a / 250
                    z = []
                    z.append(b[0: 250])
                    for i in range(1, int(c) + 1):
                        y = b[(i * 250) + 1:(i + 1) * 250 + 1]
                        z.append(y)
                    j = 3
                    for d3 in z:
                        ws_fm["D" + str(j)].value = d3
                        ws_fm["B" + str(j)].value = "-"
                        wb_fm.save(result_path)
                        j += 1
                else:
                    ws_fm['B3'].value = '-'
                    ws_fm['D3'].value = comment
            d_i = max((a.row for a in ws_fm['D'] if a.value is not None))
            wb_fm.save(result_path)
            for i in range(2, 3):
                if ws_or["R" + str(i)].value is None:
                    pass
                else:
                    ws_fm['D' + str(d_i + 1)].value = 'ManuFacturer:' + ws_or['R' + str(i)].value
                    wb_fm.save(result_path)
                if ws_or["S" + str(i)].value is None:
                    pass
                else:
                    if ws_fm['D' + str(d_i + 1)].value is None:
                        ws_fm['D' + str(d_i + 1)].value = ''
                        ws_fm['D' + str(d_i + 2)].value = ws_fm['D' + str(
                            d_i + 1)].value + ',' + 'ModelNumber:' + ws_or["S" + str(i)].value
                        wb_fm.save(result_path)
                    else:
                        ws_fm['D' + str(d_i + 1)].value = ws_fm[
                                                                    'D' + str(d_i)].value + ',' + 'ModelNumber:' + \
                                                                ws_or["S" + str(i)].value
                        wb_fm.save(result_path)
                if ws_or["W" + str(i)].value is None:
                    pass
                else:
                    if ws_fm['D' + str(d_i + 1)].value is None:
                        ws_fm['D' + str(d_i + 1)].value = ''
                        ws_fm['D' + str(d_i + 2)].value = ws_fm['D' + str(
                            d_i + 1)].value + ',' + 'DepartmentType:' + ws_or["W" + str(i)].value
                        wb_fm.save(result_path)
                    else:
                        ws_fm['D' + str(d_i + 1)].value = ws_fm['D' + str(
                            d_i)].value + ',' + 'DepartmentType:' + ws_or["W" + str(i)].value
                        wb_fm.save(result_path)
                if ws_or["U" + str(i)].value is None:
                    pass
                else:
                    if ws_fm['D' + str(d_i + 1)].value is None:
                        ws_fm['D' + str(d_i + 1)].value = ''
                        ws_fm['D' + str(d_i + 2)].value = ws_fm['D' + str(
                            d_i + 1)].value + ',' + 'SerialNumber:' + ws_or["U" + str(i)].value
                        break
                    else:
                        ws_fm['D' + str(d_i + 1)].value = ws_fm['D' + str(
                            d_i)].value + ',' + 'SerialNumber:' + ws_or["U" + str(i)].value
                        wb_fm.save(result_path)
            if ws_fm['D'+str(d_i+1)].value is not None:
                ws_fm['B'+str(d_i+1)].value ='-'
            else:
                pass
            wb_fm.save(result_path)
            e = ws_fm.i+1
            i=2
            for j in range(2, ws_or.i+1):
                if (ws_or['P' + str(i)].value == ws_or['P' + str(i + 1)].value) and (ws_or['P' + str(i)].value is not None):
                    #print('-이름이 같음')
                    e = ws_fm.i + 1
                    # 앞선 방법과 동일하게 처리
                    ws_fm["B"+str(e)].value ="-"
                    ws_fm["D"+str(e)].value = ws_or["P"+str(i)].value
                    wb_fm.save(result_path)
                    e = ws_fm.i + 1
                    for j in range(2, ws_or.i + 1):
                        e = ws_fm.i + 1
                        ws_fm["A" + str(e)].value = ws_or["A" + str(i)].value
                        ws_fm["C" + str(e)].value = ws_or["C" + str(i)].value
                        ws_fm["E" + str(e)].value = ws_or["F" + str(i)].value
                        ws_fm["F" + str(e)].value = ws_or["G" + str(i)].value
                        if len(ws_or["E" + str(i)].value) > 250:
                            a = ws_or["E" + str(i)].value
                            under = a[0:250]
                            under2 = under.replace('\n','')
                            under3 = under2.replace('\t','')
                            under4 = under3.replace('\0','')
                            under5 = under4.replace('\r','')
                            over = a[250:]
                            over2 = over.replace('\n','')
                            over3 = over2.replace('\t','')
                            over4 = over3.replace('\0','')
                            over5 = over4.replace('\r','')
                            over6 = over5.replace('\r','')
                            ws_fm["D" + str(e)].value = under5
                            ws_fm["B" + str(e)].value = "="
                            ws_fm["D" + str(e + 1)].value = over5
                            c = ws_or["J" + str(i)].value
                            if c is not None:
                                ah = c.replace('\n', ' ')
                                ah2 = ah.replace('\t', ' ')
                                if len(ah) > 250:
                                    ah_len = len(ah)
                                    bc = ah
                                    cd = ah_len / 250
                                    za = []
                                    za.append(bc[0: 250])
                                    for i in range(1, int(cd) + 1):
                                        yz = bc[(i * 250) + 1:(i + 1) * 250 + 1]
                                        za.append(yz)
                                    for dz in za:
                                        if dz == '':
                                            break
                                        d_i = max((a.row for a in ws_fm['D'] if a.value is not None))
                                        ws_fm["B" + str(d_i + 1)].value = '='
                                        ws_fm["D" + str(d_i + 1)].value = dz
                                        wb_fm.save(result_path)
                                else:
                                    ws_fm["B" + str(e + 1)].value = '='
                                    ws_fm["D" + str(e + 1)].value = ah2
                            wb_fm.save(result_path)
                            i += 1
                        else:
                            b = ws_or["E" + str(i)].value
                            if b is not None:
                                ax = b.replace('\n', ' ')
                                ax2 = ax.replace('\0', '')
                                av = ax2.replace('\r', '')
                                axc = av.replace('\t', '')
                                ws_fm["D" + str(e)].value = axc

                            c = ws_or["J" + str(i)].value
                            if c is not None:
                                ah = c.replace('\n', ' ')
                                ah2 = ah.replace('\t', ' ')
                                if len(ah) > 250:
                                    ah_len = len(ah)
                                    bc = ah
                                    cd = ah_len / 250
                                    za = []
                                    za.append(bc[0: 250])
                                    for i in range(1, int(cd) + 1):
                                        yz = bc[(i * 250) + 1:(i + 1) * 250 + 1]
                                        za.append(yz)
                                    for dz in za:
                                        if dz == '':
                                            break
                                        d_i = max((a.row for a in ws_fm['D'] if a.value is not None))
                                        ws_fm["B" + str(d_i + 1)].value = '='
                                        ws_fm["D" + str(d_i + 1)].value = dz
                                        wb_fm.save(result_path)
                                else:
                                    ws_fm["B" + str(e + 1)].value = '='
                                    ws_fm["D" + str(e + 1)].value = ah2
                            wb_fm.save(result_path)
                            i +=1
                        if ws_or["A" + str(i)].value is None:
                            break

                if (ws_or['P' + str(i)].value != ws_or['P' + str(i + 1)].value) and (ws_or['P' + str(i)].value is not None):
                        #print("- 이름이 다름")
                        e = ws_fm.i + 1
                        ws_fm["B"+str(e)].value ="-"
                        ws_fm["D"+str(e)].value =ws_or["P"+str(i)].value
                        wb_fm.save(result_path)
                        e = ws_fm.i+1
                        ws_fm["A"+str(e)].value = ws_or["A"+str(i)].value
                        ws_fm["C"+str(e)].value = ws_or["C"+str(i)].value
                        ws_fm["E" + str(e)].value = ws_or["F" + str(i)].value
                        ws_fm["F" + str(e)].value = ws_or["G" + str(i)].value
                        if len(ws_or["E"+str(i)].value) > 250:
                            a=ws_or["E"+str(i)].value
                            under = a[0:250]
                            under2 = under.replace('\n',' ')
                            under3 = under2.replace('\t','')
                            under4 = under3.replace('\0','')
                            under5 = under4.replace('\r','')
                            over = a[250:]
                            over2 = over.replace('\n',' ')
                            over3 = over2.replace('\t','')
                            over4 = over3.replace('\0','')
                            over5 = over4.replace('\r','')
                            ws_fm["D" + str(e)].value = under5
                            ws_fm["B"+ str(e+1)].value ="="
                            ws_fm["D"+str(e+1)].value =over5
                            c = ws_or["J" + str(i)].value
                            if c is not None:
                                ah = c.replace('\n', ' ')
                                ah2 = ah.replace('\t', ' ')
                                if len(ah) > 250:
                                    ah_len = len(ah)
                                    bc = ah
                                    cd = ah_len / 250
                                    za = []
                                    za.append(bc[0: 250])
                                    for i in range(1, int(cd) + 1):
                                        yz = bc[(i * 250) + 1:(i + 1) * 250 + 1]
                                        za.append(yz)
                                    for dz in za:
                                        if dz == '':
                                            break
                                        d_i = max((a.row for a in ws_fm['D'] if a.value is not None))
                                        ws_fm["B" + str(d_i + 1)].value = '='
                                        ws_fm["D" + str(d_i + 1)].value = dz
                                        wb_fm.save(result_path)
                                else:
                                    ws_fm["B" + str(e + 1)].value = '='
                                    ws_fm["D" + str(e + 1)].value = ah2
                            wb_fm.save(result_path)
                            e=ws_fm.i+1
                            i +=1

                        b = ws_or["E" + str(i)].value
                        if b is not None:
                            ax = b.replace('\n', ' ')
                            ax2 = ax.replace('\0', '')
                            av = ax2.replace('\r', '')
                            axc = av.replace('\t', '')
                            ws_fm["D" + str(e)].value = axc
                        c = ws_or["J" + str(i)].value
                        if c is not None:
                            ah = c.replace('\n', ' ')
                            ah2 = ah.replace('\t', ' ')
                            if len(ah) > 250:
                                ah_len = len(ah)
                                bc = ah
                                cd = ah_len / 250
                                za = []
                                za.append(bc[0: 250])
                                for i in range(1, int(cd) + 1):
                                    yz = bc[(i * 250) + 1:(i + 1) * 250 + 1]
                                    za.append(yz)
                                for dz in za:
                                    if dz == '':
                                        break
                                    d_i = max((a.row for a in ws_fm['D'] if a.value is not None))
                                    ws_fm["B" + str(d_i + 1)].value = '='
                                    ws_fm["D" + str(d_i + 1)].value = dz
                                    wb_fm.save(result_path)
                            else:
                                ws_fm["B" + str(e + 1)].value = '='
                                ws_fm["D" + str(e + 1)].value = ah2
                        wb_fm.save(result_path)
                        i+=1

if __name__ == "__main__":
    main()
