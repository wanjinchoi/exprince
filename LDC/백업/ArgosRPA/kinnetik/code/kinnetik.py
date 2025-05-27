import openpyxl
import pandas as pd
import os


#form_path='C:\\work\\LDC\\form.xlsx'
#excel_path='C:\\work\\LDC\\excel\\xlsx_file\\'
#result_path='C:\\work\\LDC\\k_result.xlsx'
#except_path='C:\\work\\LDC\\Name.xlsx'


form_path='C:\\ArgosRPA\\form\\form.xlsx'
excel_path='C:\\ArgosRPA\\kinnetik\\excel\\xlsx_file\\'
result_path='C:\\ArgosRPA\\kinnetik\\result\\k_result.xlsx'
except_path='C:\\ArgosRPA\\kinnetik\\form\\except.xlsx'

def kpartname2(result_path):
    xlsx_file = os.listdir(excel_path)
    for xfile in xlsx_file:
        if xfile == 'kpartname2.xlsx':
            wb = openpyxl.load_workbook(excel_path + xfile)
            ws = wb.active
            wb_fm = openpyxl.load_workbook(result_path)
            ws_fm = wb_fm.active
            j = 2
            for i in range(1, ws.i+1):
                if ws['A'+str(i)].value is None:
                    break
                else:
                    ws_fm["D"+str(j)].value = ws_fm['D'+str(j)].value +', '+ ws['A'+str(i)].value
                    a = ws_fm["D"+str(j)].value
                    b = a.replace("혻",'')
                    ws_fm["D" + str(j)].value = b
                j+=1
                wb_fm.save(result_path)



def main(excelpath):
    csv_file = os.listdir(excelpath)
    for file in csv_file:
        try:
                if file =='kitemno.csv':
                    df = pd.read_csv(excelpath + '\\' + file,encoding='cp949')
                    df.to_excel(excel_path+'kitemno.xlsx', index=False)
                if file =='kpartname.csv':
                    df = pd.read_csv(excelpath + '\\' + file,encoding='cp949',sep='\t')
                    df.to_excel(excel_path+'kpartname.xlsx', index=False)
                if file == 'kpartno.csv':
                    df = pd.read_csv(excelpath + '\\' + file,encoding='cp949',sep='\t')
                    df.to_excel(excel_path+'kpartno.xlsx', index=False)
                if file =='kqty.csv':
                    df = pd.read_csv(excelpath + '\\' + file,encoding='cp949')
                    df.to_excel(excel_path+'kqty.xlsx', index=False)
                if file =='kuom.csv':
                    df = pd.read_csv(excelpath + '\\' + file,encoding='cp949')
                    df.to_excel(excel_path+'kuom.xlsx', index=False)
                if file == 'kpartname2.csv':
                    df = pd.read_csv(excelpath + '\\' + file,encoding='cp949',sep='\t')
                    df.to_excel(excel_path + 'kpartname2.xlsx', index=False)
        except:
            file2= file.replace('.csv','.xlsx')
            wb = openpyxl.load_workbook(except_path)
            ws = wb.active
            wb.remove(ws)
            wb.create_sheet('Sheet1')
            wb.save(excel_path + file2)

    wb_fm = openpyxl.load_workbook(form_path)
    ws_fm = wb_fm.active
    xlsx_file = os.listdir(excel_path)
    for xfile in xlsx_file:
        if xfile =='kitemno.xlsx':
            wb = openpyxl.load_workbook(excel_path + xfile)
            ws = wb.active
            wb_fm = openpyxl.load_workbook(form_path)
            ws_fm = wb_fm.active
            for i in range(1,ws.i+1):
                a_i = max((a.row for a in ws_fm['A'] if a.value is not None))
                ws_fm['A' + str(a_i+1)].value = ws['A' + str(i)].value
                wb_fm.save(result_path)


        if xfile == 'kpartno.xlsx':
            wb = openpyxl.load_workbook(excel_path + xfile)
            ws = wb.active
            wb_fm = openpyxl.load_workbook(result_path)
            ws_fm = wb_fm.active
            for i in range(1, ws_fm.i + 1):
                c_row = max((a.row for a in ws_fm['C'] if a.value is not None))
                c_i = c_row + 1
                if ws['A'+str(i)].value is None:
                    break
                ws_fm['C'+str(c_i)].value = ws['A'+str(i)].value
                wb_fm.save(result_path)


        if xfile == 'kpartname.xlsx':
            wb = openpyxl.load_workbook(excel_path + xfile)
            ws = wb.active
            wb_fm = openpyxl.load_workbook(result_path)
            ws_fm = wb_fm.active

            for i in range(1, ws.i+1):
                d_row = max((a.row for a in ws_fm['D'] if a.value is not None))
                if ws['B'+str(i)].value is not None:
                    ws_fm['D'+str(d_row+1)].value = ws['A'+str(i)].value + ws['B'+str(i)].value
                    wb_fm.save(result_path)
                else:
                    ws_fm['D' + str(d_row + 1)].value = ws['A' + str(i)].value
                    wb_fm.save(result_path)

            kpartname2(result_path)


        if xfile == 'kqty.xlsx':
            wb = openpyxl.load_workbook(excel_path + xfile)
            ws = wb.active
            wb_fm = openpyxl.load_workbook(result_path)
            ws_fm = wb_fm.active
            for i in range(1, ws.i + 1):
                e_row = max((a.row for a in ws_fm['E'] if a.value is not None))
                e_i = e_row + 1
                if ws['A' + str(i)].value is None:
                    break
                ws_fm['E' + str(e_i)].value = ws['A' + str(i)].value
                wb_fm.save(result_path)

        if xfile == 'kuom.xlsx':
            wb = openpyxl.load_workbook(excel_path + xfile)
            ws = wb.active
            wb_fm = openpyxl.load_workbook(result_path)
            ws_fm = wb_fm.active
            for i in range(1, ws.i + 1):
                f_row = max((a.row for a in ws_fm['F'] if a.value is not None))
                f_i = f_row + 1
                if ws['A' + str(i)].value is None:
                    break
                ws_fm['F' + str(f_i)].value = ws['A' + str(i)].value
                wb_fm.save(result_path)








if __name__ == "__main__":
    main()

