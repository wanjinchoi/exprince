import os
import glob
import openpyxl
from datetime import datetime

today = datetime.today()
r_today = today.strftime('%Y%m%d')
file_path='C:\\ArgosRPA\\meps\\'+today.strftime('%Y%m%d')+'\\'
form_path ='C:\\ArgosRPA\\form\\form.xlsx'
result_path ='C:\\ArgosRPA\\meps\\result\\m_result.xlsx'
meps2_path = 'C:\\ArgosRPA\\meps\\excel\\meps.xlsx'

#file_path='C:\\work\\LDC\\meps\\'+r_today+'\\'
#form_path ='C:\\work\\LDC\\form.xlsx'
#result_path= 'C:\\work\\LDC\\m_result.xlsx'
#meps2_path = 'C:\\Users\\vivans\\Desktop\\test\\meps2.xlsx'

def main(excelplus):
    file = os.listdir(file_path)
    #수정날짜 순으로 정리
    flist = sorted(glob.glob(file_path + '*.xlsx'), key=os.path.getmtime)
    r_len = len(flist) - 1
    r_file = flist[r_len]
    wb = openpyxl.load_workbook(r_file)
    ws = wb.active
    wb_fm =openpyxl.load_workbook(form_path)
    ws_fm = wb_fm.active
    ws_fm['B2'].value ='+'
    wb_mp2 = openpyxl.load_workbook((meps2_path))
    ws_mp2 = wb_mp2.active


    if ws['C6'].value is None:
        ws_fm['D2'].value ='PRODUCT TYPE:' + ws['C7'].value + ', ' +'PRODUCT SERIAL:'+ws['C8'].value+', '+'MANUFACTURER:' + ws['C9'].value
    elif ws['C7'].value is None:
        ws_fm['D2'].value = 'PRODUCT BRAND:' + ws['C6'].value + ', ' +'PRODUCT SERIAL NO.'+ ws['C8'].value+', '+'MANUFACTURER:'+ws['C9'].value
    elif ws['C8'].value is None:
        ws_fm['D2'].value = 'PRODUCT BRAND:' + ws['C6'].value + ', ' + 'PRODUCT TYPE:' + ws['C7'].value + ', ' + 'MANUFACTURER:'+ws['C9'].value
    elif ws['C9'].value is None:
        ws_fm['D2'].value = 'PRODUCT BRAND:' + ws['C6'].value + ', ' + 'PRODUCT TYPE:' + ws['C7'].value + ', ' + 'PRODUCT SERIAL NO.'+ ws['C8'].value
    else:
        ws_fm['D2'].value = 'PRODUCT BRAND:' + ws['C6'].value + ', ' + 'PRODUCT TYPE:' + ws['C7'].value + ', ' + 'PRODUCT SERIAL NO.' + ws['C8'].value + ', '+'MANUFACTURER:' + ws['C9'].value
    wb_fm.save(result_path)

    j_i = max((a.row for a in ws['J'] if a.value is not None))
    e = ws_fm.i+1
    a = 1
    z = 18
    for i in range(e,j_i+1):
        d_row = max((a.row for a in ws_fm['D'] if a.value is not None))
        if ws['A'+str(z)].value is None:
            break
        ws_fm['A'+str(d_row+1)].value = ws['A'+str(z)].value
        ws_fm['C'+str(d_row+1)].value = ws['F'+str(z)].value
        if ws['D'+str(z)].value is not None:
            ws_fm['D' + str(d_row+1)].value = ws['B' + str(z)].value + '/' + ws['D' + str(z)].value            
        else:
           ws_fm['D' + str(d_row + 1)].value = ws['B' + str(z)].value 


        ws_fm["E" + str(d_row+1)].value = ws['J' + str(z)].value
        ws_fm["F" + str(d_row+1)].value = ws['K' + str(z)].value
        wb_fm.save(result_path)

        # 250글가자 넘을 경우
        excelplus = ws_fm['D'+str(i)].value
        d_row = max((a.row for a in ws_fm['D'] if a.value is not None))
        if len(excelplus) > 250:
            a = len(excelplus)
            b = excelplus
            c = a / 250
            z = []
            z.append(b[0: 250])
            for i in range(1, int(c) + 1):
                y = b[(i * 250) + 1:(i + 1) * 250 + 1]
                z.append(y)
            j = 2
            ws_fm['B2'] = '+'
            for d3 in z:
                ws_fm["D" + str(j)].value = d3
                wb_fm.save(result_path)
                j += 1
                ws_fm["B" + str(j)].value = '='
            wb_fm.save(result_path)
        else:
            pass
        if ws_mp2["A"+str(a)].value is None:
            break
        else:
            ws_fm['B'+str(d_row+1)].value = '='
            ws_fm['D'+str(d_row+1)].value = ws_mp2["A"+str(a)].value
        a +=1
        z+=1
        wb_fm.save(result_path)

if __name__ == "__main__":
    main('aa')

