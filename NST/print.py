import openpyxl
import os
import win32com.client
import math
import glob
import shutil
import datetime

class Output:

    def __init__(self, file_path, data_time):
        self.PDF_path = self.get_latest_folder(r'C:\work\NST\1.Computer\output\chengwoo')
        self.vba_path = r'C:\work\NST\python\print_sunil.bas'

        if not os.path.exists(self.PDF_path):
            os.makedirs(self.PDF_path)

        self.file_path = file_path
        self.data_time = data_time

    def file_print(self, flist):
        flist = sorted(glob.glob(self.PDF_path + '\\' + '*.xlsm'), key=os.path.getmtime)
        for i in range(len(flist)):
            try:
                wb = openpyxl.load_workbook(flist[i], keep_vba=True)
                ws = wb['5대업체 식별표서식 (사용)']
                if int(ws['BH12'].value) > 0:
                    print_quantity = math.ceil(int(ws['BH12'].value) / 2)

                    excel_app = win32com.client.DispatchEx("Excel.Application")
                    workbook = excel_app.Workbooks.Open(flist[i])

                    for _ in range(print_quantity):
                        excel_app.Run('print_sunil.print_sunil')

                    workbook.Close(SaveChanges=False)
                    excel_app.Quit()
            except Exception as e:
                print(f"Error processing file {flist[i]}: {e}")

    def start(self):
        try:
            self.file_print(self.PDF_path)
        except Exception as e:
            print(f"Error in start method: {e}")
            return 1


def do_start(file_path, data_time):
    ws = Output(file_path=file_path, data_time=data_time)
    ws.start()


def main(file_path, data_time):
    do_start(file_path, data_time)


def folder_remove():
    path = r"C:\work\NST\1.Computer\output\chengwoo"
    folders = [f for f in os.listdir(path) if os.path.isdir(os.path.join(path, f))]
    length = len(folders)
    if length > 1:
        for i in range(length):
            folder_path = os.path.join(path, folders[i])
            shutil.rmtree(folder_path)

    print("폴더삭제완료")


if __name__ == '__main__':
    file_path = r'C:\work\NST\1.Computer\standard_form\5대업체 식별표서식2023.xlsm'
    data_time = '2024-02-28'
    folder_remove()
    main(file_path, data_time)
