import openpyxl
import  glob
import  os
import pandas as pd
import datetime
import shutil
import re
import win32com.client
import math

start_ts = datetime.datetime.now().strftime('%Y%m%d-%H%M%S')

#xslm양식 파일 경로
file_path = 'C:\\work\\NST\\1.Computer\\standard_form\\5대업체 식별표서식2023.xlsm'
folder_path = rf'C:\work\NST\1.Computer\output\sunil\{start_ts}'

def folder_organize(folder_path):

    base_path = r'C:\work\NST\1.Computer\output\sunil'

    # 모든 하위 폴더 경로 가져오기
    subdirs = [f for f in glob.glob(os.path.join(base_path, '*')) if os.path.isdir(f)]

    # 수정시간 기준으로 정렬 (최신이 맨 뒤)
    sorted_subdirs = sorted(subdirs, key=os.path.getmtime)

    # 최신 폴더 제외하고 모두 삭제
    for folder in sorted_subdirs[:-1]:
        shutil.rmtree(folder)



def file_print(folder_path):
    vba_path = r'C:\work\NST\python\print_sunil.bas'
    flist = sorted(glob.glob(folder_path + '\\' + '*.xlsm'),key=os.path.getmtime)
    for i in range(len(flist)):
        wb = openpyxl.load_workbook(flist[i], keep_vba=True)
        ws = wb['5대업체 식별표서식 (사용)']
        if int(ws['BH12'].value) > 0:
            # 한 면에 서식표 2개가 있어서 한 장 인쇄시에 서식표 2장 출력됨. 프린트 개수 계산하는 변수
            print_quantity = math.ceil(int(ws['BH12'].value) / 2)

            vba_path = vba_path
            # Excel 애플리케이션 시작
            excel_app = win32com.client.DispatchEx("Excel.Application")
            # excel_app.Visible = False  # Excel 창을 보이게 하려면 True로 설정
            # 새로운 워크북 생성 또는 기존의 워크북 열기
            workbook = excel_app.Workbooks.Open(flist[i])
            # # excel 새로 고침
            # # workbook.RefreshAll()
            # # workbook.Save()
            #
            # # 매크로 실행 - 서식표 인쇄 VBA 매크로
            for _ in range(print_quantity):
                excel_app.Run('print_sunil.print_sunil')

            workbook.Close(SaveChanges=False)

    # Excel 어플리케이션 종료
    excel_app.Quit()





def append_to_excel(filtered_data, data):
    # 새로운 데이터 프레임 생성
    df = pd.DataFrame(filtered_data)
    df2 = pd.DataFrame(data)
    # 엑셀 파일 읽기
    wb = openpyxl.load_workbook('C:\\work\\NST\\1.Computer\\standard_form\\5대업체 식별표서식2023.xlsm',keep_vba=True)
    ws = wb['저장']
    # 이전 품번
    previous_pum_number = None
    for (index1, row1), (index2, row2) in zip(df.iterrows(), df2.iterrows()):
        start_row = max((a.row for a in ws['D'] if a.value is not None)) + 1
        today = datetime.datetime.now().strftime('%y/%m/%d')
        # 품명
        pum_name = row1['제품명']
        # 품번
        pum_number = row2['제품코드']
        # 규격
        standard = row1['제품규격']
        ###lot 번호
        lot_num = row2['업체 로트 번호']
        if pum_number == previous_pum_number:
            # 품번이 중복된 경우, I열에만 lot 번호 추가
            existing_value = ws['I' + str(start_row - 1)].value
            ws['I' + str(start_row - 1)].value = existing_value + '\n' + lot_num
            # I열의 개수를 계산하여 H열에 입력
            lot_count = len(ws['I' + str(start_row - 1)].value.split('\n'))
            ws['H' + str(start_row - 1)].value = lot_count
        else:
            # 품번이 중복되지 않은 경우, 새로운 행에 입력
            ws['A' + str(start_row)].value = today
            ws['B' + str(start_row)].value = '선일'
            # 품명
            ws['C' + str(start_row)].value = pum_name
            # 품번
            ws['D' + str(start_row)].value = pum_number
            # 규격
            ws['E' + str(start_row)].value = standard
            # lot 번호
            ws['I' + str(start_row)].value = lot_num
            # I열의 개수를 계산하여 H열에 입력 (새로 입력된 경우는 항상 1개)
            ws['F' + str(start_row)].value = row1['표면\n처리사양']
            ws['H' + str(start_row)].value = 1

        # 현재 품번을 이전 품번으로 저장
        previous_pum_number = pum_number

    # 변경 사항 저장
    try:
        wb.save('C:\\work\\NST\\1.Computer\\standard_form\\5대업체 식별표서식2023.xlsm')
    except Exception as e:
        print(f"Error occurred: {e}")



def main(checklist):
    #폴더내 엑셀파일 불러오기
    file_list = sorted(glob.glob('C:\\work\\NST\\선일\\' + '*.xlsx'), key=os.path.getmtime)
    #파일 명 가져오기
    data_file = file_list[0]

    # 출력 폴더 생성
    if not os.path.exists(folder_path):
        os.makedirs(folder_path)

    #엑셀 읽기
    data_list = pd.read_excel(data_file, engine='openpyxl', header=0)

    #가져올 엑셀 열 리스트
    fields_to_extract = ['LOT번호','전 철통','제품코드', '제품명','제품규격','발주중량']

    #합친 데이터 들어가는 리스트
    all_data=[]
    # 업체 LOT 정보에 들어가는 필드 데이터 합치기
    for index, row in data_list.iterrows():
        if pd.isna(row['제품명']):
            continue
        else:
            value = row['발주중량']
            if pd.notna(value):
                weight = int(value)
            else:
                weight = 0
            # 3가지 필드 합침
            com_lot = f"{int(row['LOT번호'])}    {int(row['전 철통'])}    {weight}"

            #합친데이터
            ex_data = [row[field] for field in fields_to_extract]
            ex_data.append(com_lot)
            all_data .append(ex_data)

    #datafram 생성
    df_result = pd.DataFrame(all_data,columns=fields_to_extract + ['업체 로트 번호'])
    #제품코드가 NAN인거 제거
    df_result = df_result[df_result['제품코드'].notna()].copy()

    # 'NST LOT No.' 컬럼 생성
    if '업체 로트 번호' in df_result.columns:
        df_result['NST LOT No.'] = (df_result.groupby('제품코드')['업체 로트 번호'].transform(lambda x: (x.str.split().str[0] != x.str.split().str[0].shift()).cumsum()))
        grouped_data = df_result[['제품코드', '업체 로트 번호', 'NST LOT No.']]
    else:
        # '업체 로트 번호'가 없을 경우 처리
        grouped_data = df_result[['제품코드']]

    #식별표 서식에 데이터 넣기
    for index, row in grouped_data.iterrows():
        # 파일 복사
        copy_path = folder_path + f'\\{index}.xlsm'
        flist = sorted(glob.glob(folder_path + '\\' + '*.xlsm'),key=os.path.getmtime)
        if len(flist) > 0:
            previous_path = flist[len(flist) - 1]
        # 이전 파일 품번과 비교하기 위해

        # 처음 한번은 입력
        if index == 0:
            shutil.copy(file_path, copy_path)
            wb = openpyxl.load_workbook(copy_path, keep_vba=True)
            ws = wb['5대업체 식별표서식 (사용)']
            # 품번입력
            non_pattern = re.compile(r'\D')
            if ws['BI12'].value and '\n' in ws['BI12'].value:
                ws['BI12'].value = None
                # ws['AI23'].value = None
            if non_pattern.search(str(row['제품코드'])):
                ws['BE12'].value = row['제품코드']
            else:
                ws['BE12'].value = int(row['제품코드'])

            # 데이터 회사명
            ws['BD12'] = '선일'

            # 그룹으로 묶여있는 업체 로트 번호 가져옴
            row_values = row['업체 로트 번호']
            # 업체 로트 번호 개수 저장
            # lot_num = row['lot_num']

            # ws['BE13'] = int(row['전 Lot 정보 - 품번'])
            # 철통 수량 - 업체 로트 번호의 개수
            # ws['BH12'].value = lot_num

            # 포트번호에 하나씩 쌓기
            if (ws['BI12'].value is not None):
                ws['BI12'].value = ws['BI12'].value + '\n' + row_values
                # ws['AI23'].value = ws['AI23'].value + '\n' + row_values
            else:
                ws['BI12'].value = row_values
                # ws['AI23'].value = row_values
            ws['BG12'].value = row['NST LOT No.']
            ws['BH12'].value = '1'
            wb.save(copy_path)

        else:
            # 이전 입력 품번이랑 같은지 판단
            previous_value = grouped_data.at[index - 1, '제품코드']
            previous_lot_no = grouped_data.at[index - 1, 'NST LOT No.']
            # 조건 1 품번이 같냐?
            if (str(row['제품코드']) == str(previous_value)) and (str(previous_lot_no) == str(row['NST LOT No.'])):
                wb = openpyxl.load_workbook(previous_path, keep_vba=True)
                ws = wb['5대업체 식별표서식 (사용)']
                row_values = row['업체 로트 번호']
                # 로트번호 입력
                ws['BI12'].value = ws['BI12'].value + '\n' + row_values
                # ws['AI23'].value = ws['AI23'].value + '\n' + row_values
                # 로트번호 수 파악
                lines = ws['BI12'].value.split('\n')
                # 철통수
                ws['BH12'].value = str(len(lines))
                wb.save(previous_path)
                wb.close()
            # 이전 파일과 품번 및 LOT넘버가 같지 않을 때
            else:
                shutil.copy(file_path, copy_path)
                wb = openpyxl.load_workbook(copy_path, keep_vba=True)
                ws = wb['5대업체 식별표서식 (사용)']
                previous_value = grouped_data.at[index - 1, '제품코드']
                previous_lot_no = grouped_data.at[index - 1, 'NST LOT No.']
                non_pattern = re.compile(r'\D')
                # 품번입력
                if non_pattern.search(str(row['제품코드'])):
                    ws['BE12'].value = row['제품코드']
                else:
                    ws['BE12'].value = int(row['제품코드'])
                #
                ws['BD12'] = '선일'
                row_values = row['업체 로트 번호']
                ws['BI12'].value = row_values
                # ws['AI23'].value = row_values
                ws['BG12'].value = row['NST LOT No.']
                # 포트번호수
                lines = ws['BI12'].value.split('\n')
                # 철통수
                ws['BH12'].value = str(len(lines))
                wb.save(copy_path)
    end_dir = r'C:\work\NST\선일\endfile'
    shutil.move(data_file, os.path.join(end_dir, os.path.basename(data_file)))
    #프린트
    file_print(folder_path)
    # #양식에 해당 내용넣기
    append_to_excel(data_list, grouped_data)
    #폴더삭제
    folder_organize(folder_path)

if __name__ == "__main__":
    main('aa')

