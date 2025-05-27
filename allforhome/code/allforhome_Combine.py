import os

import openpyxl


def main(path):
    possible_xlsx_extension = ['.xlsx']
    # 붙여넣을 sheet명
    # ws_w = wb_w.create_sheet("Sheet1")
    for district_dir in os.listdir(path):  # root 폴더 내 권역 폴더 확인
        district_path = os.path.join(path, district_dir)  # root+권역내 폴더명(join으로 경로명을 합쳐준다.)
        for turn_dir in os.listdir(district_path):# 권역 폴더 내 회차 폴더 확인
            cnt = 7  # 복사 row범위(초기화방지를 위해 for문 밖에서)
            wb_w = openpyxl.load_workbook(r'C:\Users\vivans\Desktop\올포홈\서식\RPA_Combine.xlsx', data_only=False)  # 불러올 서식파일
            ws_w = wb_w["권역별합계"]  # 불러올 sheet명
            turn_path = os.path.join(district_path, turn_dir)
            if "RPA_result.xlsx" not in os.listdir(turn_path):
                for file in os.listdir(turn_path):  # 회차 폴더 내 파일 확인
                    file_path = os.path.join(turn_path, file)
                    print(file_path)
                    if ('(올포홈)거래명세서.xlsx' in file_path) or ('RPA_result.xlsx' in file_path):
                        pass
                    elif 'RPA_result' not in file_path:
                        ext = os.path.splitext(file_path)[1]
                        print(ext)
                        if ext in possible_xlsx_extension:  # 엑셀 파일만 작업
                            wb_r = openpyxl.load_workbook(file_path, data_only=True)  # 불러올 파일
                            ws_r = wb_r["내역서"]  # 가져올 sheet명
                            for i in range(7, ws_r.i + 1):  # 가져올 범위 설정
                                cnt_a = 1
                                for j in range(1, ws_r.max_column):
                                    # 붙여넣을 범위 =  가져올범위
                                    ws_w.cell(row=cnt, column=cnt_a).value = ws_r.cell(row=i, column=j).value
                                    cnt_a += 1
                                cnt += 1
                    else:
                        pass

                    wb_w.save(turn_path + r'\RPA_result.xlsx')  # 파일 저장




if __name__ == "__main__":
    main(r'C:\Users\vivans\Desktop\올포홈\allforhome')
