import os

import openpyxl


def main(path):
    for district_dir in os.listdir(path):  # root 폴더 내 권역 폴더 확인
        print(district_dir)
        district_path = os.path.join(path, district_dir)  # root+권역내 폴더명(join으로 경로명을 합쳐준다.)
        for turn_dir in os.listdir(district_path):  # 권역 폴더 내 회차 폴더 확인
            turn_path = os.path.join(district_path, turn_dir)
            k = list(os.listdir(turn_path))
            print(k)
            for file in os.listdir(turn_path):  # 회차 폴더 내 파일 확인
                file_path = os.path.join(turn_path, file)
                if ('RPA_result.xlsx' in k) and ('(올포홈)거래명세서.xlsx' in k):
                        wb = openpyxl.load_workbook(turn_path + r'\RPA_result.xlsx', data_only=True)
                        ws = openpyxl.load_workbook(turn_path + r'\(올포홈)거래명세서.xlsx', data_only=True)
                        wb_r = wb["건설,장기수선"]  # 1
                        ws_r = ws["거래명세서"]  # 2
                        wa_r = wb["그외코드"]  # 3

                        # 두 시트 비교해서 row의 최대길이 가져오기
                        a = wb_r.i
                        b = wa_r.i
                        if a > b:
                            row_len = a
                        else:
                            row_len = b
                        for i in range(2, row_len + 1):
                            for j in range(11, ws_r.i + 1):
                                code = wb_r["A" + str(i)].value
                                code2 = ws_r["A" + str(j)].value
                                code3 = wa_r["A" + str(i)].value
                                if code == code2:
                                    try:
                                        ws_r["E" + str(j)].value = wb_r["I" + str(i)].value
                                    except AttributeError as err:
                                        break
                                elif code3 == code2:
                                    try:
                                        ws_r["E" + str(j)].value = wa_r["I" + str(i)].value
                                    except AttributeError as err:
                                        break
                                else:
                                    print("e")
                        ws.save(turn_path + r'\(올포홈)거래명세서.xlsx')
                else:
                     pass


if __name__ == '__main__':
    main(r'C:\Users\vivans\Desktop\올포홈\allforhome')
