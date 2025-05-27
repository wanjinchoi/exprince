import os

import openpyxl


def set_col(ow):
    ow["A1"].value = "단가코드"
    ow["B1"].value = "명칭"
    ow["C1"].value = "규격"
    ow["D1"].value = "단위"
    ow["E1"].value = "단가/자재비"
    ow["F1"].value = "단가/인건비"
    ow["G1"].value = "단가/경비"
    ow["H1"].value = "단가/합계"
    ow["I1"].value = "수량"
    ow["J1"].value = "단위"
    ow["K1"].value = "계약단가/자재비"
    ow["L1"].value = "계약단가/인건비"
    ow["M1"].value = "계약단가/경비"
    ow["N1"].value = "계약단가/합계"
    ow["O1"].value = "산출금액/자재비"
    ow["P1"].value = "산출금액/인건비"
    ow["Q1"].value = "산출금액/경비"
    ow["R1"].value = "산출금액/합계"
    ow["S1"].value = "권역구분"


def set_col2(ow2):
    ow2["A1"].value = "단가코드"
    ow2["B1"].value = "명칭"
    ow2["C1"].value = "규격"
    ow2["D1"].value = "단위"
    ow2["E1"].value = "단가/자재비"
    ow2["F1"].value = "단가/인건비"
    ow2["G1"].value = "단가/경비"
    ow2["H1"].value = "단가/합계"
    ow2["I1"].value = "수량"
    ow2["J1"].value = "단위"
    ow2["K1"].value = "계약단가/자재비"
    ow2["L1"].value = "계약단가/인건비"
    ow2["M1"].value = "계약단가/경비"
    ow2["N1"].value = "계약단가/합계"
    ow2["O1"].value = "산출금액/자재비"
    ow2["P1"].value = "산출금액/인건비"
    ow2["Q1"].value = "산출금액/경비"
    ow2["R1"].value = "산출금액/합계"
    ow2["S1"].value = "권역구분"


def main(path):
    for district_dir in os.listdir(path):  # root 폴더 내 권역 폴더 확인
        # district_dir_lower = district_dir.lower()
        district_path = os.path.join(path, district_dir)  # root+권역내 폴더명(join으로 경로명을 합쳐준다.)
        for turn_dir in os.listdir(district_path):
            turn_path = os.path.join(district_path, turn_dir)  # 권역 폴더 내 회차 폴더 확인
            for file in os.listdir(turn_path):  # 회차 폴더 내 파일 확인
                file_path = os.path.join(turn_path, file)
                print(file_path)
                if '(올포홈)거래명세서.xlsx' in file_path:
                    pass
                elif 'RPA_result' in file_path:
                    code_list = []
                    else_list = []
                    wb = openpyxl.load_workbook(file_path, data_only=True)
                    sh = wb["권역별합계"]
                    sheet_name_list = wb.sheetnames
                    if ("건설,장기수선" not in sheet_name_list) or ("그외코드" not in sheet_name_list):
                        ow = wb.create_sheet("건설,장기수선")
                        ow2 = wb.create_sheet("그외코드")
                        set_col(ow)
                        set_col(ow2)
                        for i in range(7, sh.i + 1):
                            code = sh["H" + str(i)].value
                            name = sh["I" + str(i)].value
                            standard = sh["J" + str(i)].value
                            unit = sh["K" + str(i)].value
                            dan_cost = sh["L" + str(i)].value
                            dan_labor = sh["M" + str(i)].value
                            dan_budget = sh["N" + str(i)].value
                            dan_sumit = sh["O" + str(i)].value
                            amount = sh["P" + str(i)].value
                            unit2 = sh["Q" + str(i)].value
                            con_cost = sh["R" + str(i)].value
                            con_labor = sh["S" + str(i)].value
                            con_budget = sh["T" + str(i)].value
                            con_sumit = sh["U" + str(i)].value
                            rls_cost = sh["V" + str(i)].value
                            rls_labor = sh["W" + str(i)].value
                            rls_budget = sh["X" + str(i)].value
                            rls_sumit = sh["Y" + str(i)].value
                            ere = sh["BL" + str(i)].value
                            try:
                                if "건설" in ere:
                                    print("rr")
                                    if code in code_list:
                                        print("찾아서..amount만 더하고")
                                        for j in range(2, ow.i + 1):
                                            code_r = ow["A" + str(j)].value
                                            if code == code_r:
                                                ow["I" + str(j)].value = ow["I" + str(j)].value + amount
                                    else:
                                        code_list.append(code)

                                        list_to_append = [code, name, standard, unit, dan_cost, dan_labor, dan_budget,
                                                          dan_sumit, amount, unit2, con_cost, con_labor, con_budget,
                                                          con_sumit, rls_cost, rls_labor, rls_budget, rls_sumit, ere]
                                        ow.append(list_to_append)
                                else:
                                    if code in else_list:
                                        for k in range(2, ow2.i+1):
                                            code_a = ow2["A"+str(k)].value
                                            if code == code_a:
                                                ow2["I" + str(k)].value = ow2["I" + str(k)].value + amount

                                    else:
                                        else_list.append(code)

                                        list_to_else = [code, name, standard, unit, dan_cost, dan_labor, dan_budget, dan_sumit,amount, unit2, con_cost, con_labor, con_budget, con_sumit, rls_cost,rls_labor, rls_budget, rls_sumit, ere]
                                        ow2.append(list_to_else)

                            except TypeError as e:
                                print(e)
                    else:
                        pass
                    wb.save(file_path)

if __name__ == '__main__':
    main(r'C:\Users\vivans\Desktop\올포홈\allforhome')
