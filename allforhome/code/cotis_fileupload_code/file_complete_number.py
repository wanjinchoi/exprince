import openpyxl
import os
import glob

from datetime import datetime
from dateutil.relativedelta import relativedelta



def main(index,folder_num):
    # original_file_path = 'X:\\〔1〕   완 료 신 청 = A\\'
    ##파일경로
    original_file_path = 'C:\\〔1〕   완 료 신 청 = 첨부 B\\'
    # 폴더이름 가져오기
    folders = [f for f in os.listdir(original_file_path) if os.path.isdir(os.path.join(original_file_path, f))]
    folder_name = folders[int(folder_num)]
    # 폴더경로
    folder_path = original_file_path + folder_name + '\\'
    # 엑셀 파일이름 가져오기
    flist = sorted(glob.glob(folder_path + '*.xlsx'), key=os.path.getmtime)
    excel_path = flist[0]

    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active
    number = ws['A'+str(index)].value
    max_row = max((a.row for a in ws['A'] if a.value is not None))
    number_length = max_row
    if int(index) > int(number_length):
        result ='no'
        print(result)
    else:
        print(number)

if __name__ == "__main__":
    main('4','2')

