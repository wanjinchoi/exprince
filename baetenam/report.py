import glob
import os
import openpyxl
from datetime import datetime, timedelta

def main(checklist):
    current_date_str = datetime.now().strftime("%Y-%m-%d")
    current_date = datetime.strptime(current_date_str, "%Y-%m-%d")

    yesterday_date = current_date - timedelta(days=1)
    yesterday_date_str = yesterday_date.strftime("%Y-%m-%d")
    # #store폴더경로
    folder_path = 'Z:\\viet\\store_screens\\'+yesterday_date_str+'\\'
    #report양식 경로
    form_path = 'Z:\\viet\\baetenam\\form.xlsx'
    #report저장경로
    report_path = 'Z:\\viet\report\\'
    #report이름
    report_xlsx_path='Z:\\viet\\report\\'+yesterday_date_str+'report'+'.xlsx'


    #리포트 파일이 있는지 없는지 검토
    if os.path.exists(report_xlsx_path+yesterday_date_str+'.xlsx'):
        pass
    else:
        wb = openpyxl.load_workbook(form_path)
        ws = wb.active
        wb.save(report_xlsx_path)
        wb.close()



    flist = sorted(glob.glob(folder_path + '*.png'), key=os.path.getmtime)
    r_len = len(flist) - 1
    # 파일 이름

    for i in range(0, r_len + 1):
        r_file = flist[i]
        # 파일이름 가져오기 위해서 나누기
        a = r_file.split('\\')
        file_name = a[4]
        # 파일이름 쪼개기
        a = file_name.split('_')
        #발견시간
        x = a[0]
        hours = x[:2]
        minutes = x[2:4]
        seconds = x[4:]
        detect_time = f"{hours}:{minutes}:{seconds}"
        #발견PC
        pc_name = a[1]
        #위치
        postion=a[2]
        wb = openpyxl.load_workbook(report_xlsx_path)
        ws = wb['row_data']
        column_contents = []
        for cell in ws['B']:
            if cell.value is None:
                pass
            else:
                column_contents.append(cell.value)
        i = max((a.row for a in ws['B'] if a.value is not None))
        if detect_time in column_contents:
            pass
        else:
            ws['B'+str(i+1)].value = detect_time
            ws['C'+str(i+1)].value = pc_name
            ws['D'+str(i+1)].value = postion
            wb.save(report_xlsx_path)
        wb.close()
if __name__ == "__main__":
    main('aa')
