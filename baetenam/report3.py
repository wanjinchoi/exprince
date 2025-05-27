import glob
import os
import openpyxl
from datetime import datetime, timedelta
from collections import Counter
from openpyxl.drawing.image import Image
import random







def find_matching_files(folder_path, pc_name,camera_name):
    all_files = os.listdir(folder_path)

    # Dictionary to store grouped files
    grouped_files = {}
    camera_name2 = '_'+camera_name
    # 'PC1_af'와 'PC1_bf'를 가진 파일들을 필터링
    for file in all_files:
        if pc_name in file and camera_name2 in file and ('_af' in file or '_bf' in file):
            base_name = file.replace('_af', '').replace('_bf', '')
            if base_name not in grouped_files:
                grouped_files[base_name] = [file]
            else:
                grouped_files[base_name].append(file)

    # Print or process the grouped files
    x = list(grouped_files.values())
    # Print or process the randomly selected files
    random_element = random.choice(x)
    af_file = random_element[0]
    bf_file = random_element[1]
    return af_file, bf_file

def main(checklist):
    current_date_str = datetime.now().strftime("%Y-%m-%d")
    current_date = datetime.strptime(current_date_str, "%Y-%m-%d")

    yesterday_date = current_date - timedelta(days=1)
    yesterday_date_str = yesterday_date.strftime("%Y-%m-%d")

    #store폴더경로
    folder_path = 'Z:\\viet\\store_screens\\'+yesterday_date_str+'\\'
    #report양식 경로
    form_path = 'Z:\\viet\\baetenam\\form.xlsx'
    #report저장경로
    report_path = 'Z:\\viet\report\\'
    #report이름
    report_xlsx_path='Z:\\viet\\report\\'+yesterday_date_str+'report'+'.xlsx'
    # 저장이름
    sotre_xlsx_path = 'Z:\\viet\\report\\' + yesterday_date_str + '_report' + '.xlsx'
    wb = openpyxl.load_workbook(report_xlsx_path)



    # 한꼭지(PC이름 추출하는)
    ws_d = wb['row_data']
    ws = wb['total']

    camera_counter = Counter()
    #PC랑 카메라별 몇개인지 카운팅
    for row in ws_d.iter_rows(min_row=5, values_only=True):
        pc_name, camera = row[2], row[3]
        camera_counter[(pc_name, camera)] += 1

    #카운트후 매칭하여 리스트에 담기
    x= []
    for (pc_name, camera), count in camera_counter.items():
        x.append([pc_name, camera, count])

    for data in x:
        i = max((a.row for a in ws['B'] if a.value is not None))
        if data[0] is None:
            break
        pc_name = data[0]
        camera_name = data[1]
        count = data[2]
        af_file, bf_file = find_matching_files(folder_path,pc_name,camera_name)
        a_img = Image(folder_path+af_file)
        b_img = Image(folder_path+bf_file)

        #셀높이 고정
        ws.row_dimensions[i+1].height=135.75
        cell_style = ws['B5'].style

        ws['B'+str(i+1)].value = pc_name
        # ws['B' + str(i + 1)].style = cell_style
        ws['C' + str(i + 1)].value = camera_name
        # ws['C' + str(i + 1)].style = cell_style
        ws['D' + str(i + 1)].value = count
        # ws['D' + str(i + 1)].style = cell_style

        # 이미지 추가
        ws.add_image(a_img,'G'+str(i+1))
        ws.add_image(b_img,'E'+str(i+1))
        wb.save(sotre_xlsx_path)

    if os.path.exists(report_xlsx_path):
        os.remove(report_xlsx_path)
    # wb_r = openpyxl.load_workbook(sotre_xlsx_path)
    # ws_r = wb_r['total']
    #
    # r_i = max((a.row for a in ws_r['B'] if a.value is not None))
    # sum_total =[]
    # pc_name = []
    # q=0
    # for i in range(5,r_i+1):
    #     if(ws_r['B'+str(i)].value == ws_r['B'+str(i+1)].value) and (ws_r["B"+str(i)].value is not None):
    #             sum_total.append(ws_r['D'+str(i)].value)
    #             if ws_r['B' + str(i)].value not in pc_name:
    #                 pc_name.append(ws_r['B' + str(i)].value)
    #
    #     else:
    #         k_i = max((a.row for a in ws_r['K'] if a.value is not None))
    #         total= sum(sum_total)
    #         ws_r['L'+str(k_i+1)].value = total
    #         ws_r['K'+str(k_i+1)].value = pc_name[0]
    #         pc_name=[]
    #         sum_total=[]
    #         wb_r.save(sotre_xlsx_path)
    #         wb_r.close()


if __name__ == "__main__":
    main('aa')