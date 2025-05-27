import openpyxl
import pandas as pd
file_path = 'C:\\Users\\vivans\\Desktop\\costco\\'
def main(file):
    a = open(file_path + file, 'r', encoding="UTF-8")
    l = a.readlines()
    l2=[]
    l3=[]
    for a in l:
        l2.append(a.replace(',',''))
    for b in l2:
        l3.append(b.strip())

    print(l3)
    for i in range(len(l3)):
        if '판매' not in l3[i]:
                pass
        else:
            if '판매' in l3[i]:
                i+1
                x = l3[i].replace('T', '')
                x = x.replace(' ', '')
                if i == 0:
                    wb = openpyxl.load_workbook(file_path + 'format.xlsx')
                    ws = wb.active
                else:
                    wb = openpyxl.load_workbook(file_path + 'result.xlsx')
                    ws = wb.active

                if '*' in x:
                    pass
                elif "Sub-" in x:
                    pass
                elif x == '':
                    pass
                else:
                    try:
                        int(x)
                        print("숫자")
                        z = l3[i].split(' ')
                        b_i = max((a.row for a in ws['B'] if a.value is not None))
                        ws["B" + str(a_i + 1)].value = z[1]
                        ws["C" + str(a_i + 1)].value = z[2]
                        if 'T' in z[3]:
                            c = z[3].replace('T', '')
                            ws["D" + str(a_i + 1)].value = c
                        else:
                            ws["D" + str(a_i + 1)].value = z[3]
                        wb.save(file_path + 'result.xlsx')
                    except:
                        x = l3[i].replace(' ', '')
                        a_i = max((a.row for a in ws['A'] if a.value is not None))
                        ws["A" + str(a_i + 1)].value = l3[i]
                        wb.save(file_path + 'result.xlsx')
                        print("품목")

            else:
                x = l3[i].replace('T','')
                x = x.replace(' ', '')
                if i == 0:
                    wb = openpyxl.load_workbook(file_path + 'format.xlsx')
                    ws = wb.active
                else:
                    wb = openpyxl.load_workbook(file_path + 'result.xlsx')
                    ws = wb.active

                if '*' in x:
                    pass
                elif "Sub-" in x:
                    pass
                elif x == '':
                    pass
                else:
                    try:
                        int(x)
                        print("숫자")
                        z = l3[i].split(' ')
                        b_i = max((a.row for a in ws['B'] if a.value is not None))
                        ws["B" + str(a_i + 1)].value = z[1]
                        ws["C"+str(a_i+1)].value = z[2]
                        if 'T' in z[3]:
                            c = z[3].replace('T','')
                            ws["D" + str(a_i + 1)].value = c
                        else:
                            ws["D"+str(a_i+1)].value = z[3]
                        wb.save(file_path+'result.xlsx')
                    except:
                        x = l3[i].replace(' ', '')
                        a_i = max((a.row for a in ws['A'] if a.value is not None))
                        ws["A"+str(a_i+1)].value = l3[i]
                        wb.save(file_path+'result.xlsx')
                        print("품목")




if __name__ == "__main__":
    main('costco1.txt')
