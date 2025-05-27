import pandas as pd
from openpyxl import load_workbook
import glob
import os

def main(path):
    #excel_path2 ='C:\\work\\nsu\\'
    xlsx_path = 'C:\\ARGOS RPA\\작업폴더\\발령\\'
    flist = sorted(glob.glob(xlsx_path + '*.xlsx'), key=os.path.getmtime)
    #엑셀 선택
    path = flist[0]
    excel_path = path
    workbook = load_workbook(excel_path)
    sheet_names = workbook.sheetnames
    sheet_name = sheet_names[0]
    if sheet_name == "전임-신규임용":
        #소속명 띄어쓰기 없애기
        df = pd.read_excel(excel_path, sheet_name,dtype={'사번': str})
        df["소속명"] = df['소속명'].astype(str)
        df["소속명"] = df["소속명"].str.replace(' ','')
        df["사번"] = df["사번"].apply(lambda x: str(x).zfill(6) if str(x) != "-"  else str(x))
        workbook = load_workbook(excel_path)
        sheet = workbook[sheet_name]
        start_row = 2
        start_column = 1
        for row_index, row in enumerate(df.values, start=start_row):
            for col_index, value in enumerate(row, start=start_column):
                sheet.cell(row=row_index, column=col_index, value=value)
        workbook.save(excel_path)
    elif sheet_name == "전임-재임용":
        #소속명 띄어쓰기 없애기
        df = pd.read_excel(excel_path, sheet_name,dtype={'사번': str})
        df["소속명"] = df['소속명'].astype(str)
        df["소속명"] = df["소속명"].str.replace(' ','')
        df["사번"] = df["사번"].apply(lambda x: str(x).zfill(6) if str(x) != "-"  else str(x))
        workbook = load_workbook(excel_path)
        sheet = workbook[sheet_name]
        start_row = 2
        start_column = 1
        for row_index, row in enumerate(df.values, start=start_row):
            for col_index, value in enumerate(row, start=start_column):
                sheet.cell(row=row_index, column=col_index, value=value)
        workbook.save(excel_path)
    else:
        sheet_name = "보직발령"
        #소속명 띄어쓰기 없애기
        df = pd.read_excel(excel_path, sheet_name,dtype={'사번': str})
        df["소속명"] = df['소속명'].astype(str)
        df["소속명"] = df["소속명"].str.replace(' ','')
        df["사번"] = df["사번"].apply(lambda x: str(x).zfill(6) if str(x) != "-"  else str(x))
        workbook = load_workbook(excel_path)
        sheet = workbook[sheet_name]
        start_row = 2
        start_column = 1
        for row_index, row in enumerate(df.values, start=start_row):
            for col_index, value in enumerate(row, start=start_column):
                sheet.cell(row=row_index, column=col_index, value=value)
        workbook.save(excel_path)

if __name__ == '__main__':
    main(r'C:\work\nsu\남서울대학교_발령.xlsx')