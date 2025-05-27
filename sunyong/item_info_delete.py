import openpyxl as op



def main(xlsx):
    file_path = 'C:\\ARGOS RPA\\senario1\\POS\\excel\\'
    wb = op.load_workbook(xlsx)
    ws = wb.active
    i = max((a.row for a in ws['A'] if a.value is not None))
    for i in range(1,i+1):
        if ws['A'+str(i)].value is not None and 'fo' in str(ws["A"+str(i)].value):
            ws.delete_rows(i)
            wb.save(r'C:\work\sunyong\result.xlsx')
            i-=1


if __name__ == "__main__":
    main(r'C:\work\sunyong\S SPLD RFQ2024040064.xlsx')

