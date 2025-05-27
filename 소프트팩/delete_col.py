from openpyxl import load_workbook


def main(input_path):
    wb = load_workbook(input_path)
    ws = wb.active
    for i in range(len(ws['AD'])):
        if ws['AD'+str(i+2)].value == '요청':
            ws['AB' + str(i + 2)].value = None
    wb.save(input_path)


if __name__ == '__main__':
    main(r'C:\work\소프트팩\0322_08시원본.xlsx')