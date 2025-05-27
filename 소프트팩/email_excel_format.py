from openpyxl import load_workbook
from openpyxl import Workbook


def main():
    new_wb = Workbook()
    new_ws = new_wb.active
    width_list = [10, 10, 10, 20, 10, 25, 35, 35, 35, 10, 50, 40, 10, 10, 30,
                  30, 30, 30]  # 컬럼 너비 초기화
    for i in range(len(width_list)):
        new_ws.column_dimensions[chr(65 + i)].width = width_list[i]  # 컬럼 너비 적용
    new_ws['A1'].value = "주문번호(30)"
    new_ws['B1'].value = "보내는분(40)"
    new_ws['C1'].value = "전화번호(14)"
    new_ws['D1'].value = "핸드폰번호(14)"
    new_ws['E1'].value = "우편번호(6)"
    new_ws['F1'].value = "동이하주소(100)"
    new_ws['G1'].value = "받는분(40)"
    new_ws['H1'].value = "전화번호(14)"
    new_ws['I1'].value = "핸드폰번호(14)"
    new_ws['J1'].value = "우편번호(6)"
    new_ws['K1'].value = "주소(100)"
    new_ws['L1'].value = "상품명(300)"
    new_ws['M1'].value = "상품옵션(200)"
    new_ws['N1'].value = "수량(3)"
    new_ws['O1'].value = "배송메세지(100)"
    new_ws['P1'].value = "정산구분(선,착,신)"
    new_ws['Q1'].value = "택배운임(6)"
    new_ws['R1'].value = "운송장번호(10)"

    wb = load_workbook(r"C:\Users\다현짱\Downloads\결과물\발주서목록-이메일.xlsx")
    ws = wb['Sheet']
    last_row = ws.i
    new_ws.row_dimensions[1].height = 35
    for i in range(last_row - 1):
        new_ws.row_dimensions[i+2].height = 40
        new_ws['B' + str(i + 2)].value = "소프트팩㈜"
        new_ws['C' + str(i + 2)].value = "1688-3188"
        new_ws['G' + str(i + 2)].value = ws['B' + str(i + 2)].value.split("/")[
            0]
        new_ws['H' + str(i + 2)].value = ws['H' + str(i + 2)].value
        new_ws['I' + str(i + 2)].value = ws['B' + str(i + 2)].value.split("/")[
            1]
        new_ws['K' + str(i + 2)].value = ws['C' + str(i + 2)].value
        new_ws['L' + str(i + 2)].value = ws['D' + str(i + 2)].value
        new_ws['N' + str(i + 2)].value = ws['G' + str(i + 2)].value
    new_wb.save(r"C:\Users\다현짱\Downloads\결과물\로엔그린-택배요청서.xlsx")


if __name__ == '__main__':
    main()
