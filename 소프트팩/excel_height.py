from openpyxl import load_workbook
import math


def main(input_file):
    wb = load_workbook(input_file)
    ws = wb['밸브']
    for i in range(len(ws['A'])):
        if ws['F' + str(i + 2)].value is not None:
            if math.ceil(len(list(ws['F' + str(i + 2)].value)) / 21) * 24 > 49.5:
                ws.row_dimensions[i + 2].height = math.ceil(len(list(ws['F' + str(i + 2)].value)) / 21) * 24
    wb.save(input_file)


if __name__ == '__main__':
    main()