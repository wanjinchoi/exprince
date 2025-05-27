from openpyxl.styles.borders import Border, Side
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
import time


def main(file_path):
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'),
                         top=Side(style='thin'),
                         bottom=Side(style='thin'))  # 윤각선 속성 설정
    wb = load_workbook(file_path)
    width_list = [8.38, 28.63, 54.88, 53, 13.5, 15.88]
    fontStyle = Font(size="15")  # 폰트 사이즈 초기화
    sheet_arr = ['미부착-봉투', '미부착-악세']
    for t in sheet_arr:
        wb[t].insert_cols(1)
        for i in range(len(wb[t]['B'])-1):
            wb[t]['A'+str(i+2)].value = i+1
        a = time.strftime("%H")
        if int(a) > 12:
            a = int(a) - 12
        wb[t]['A'+str(1)].value = str(a) + "시"
        for i in range(len(width_list)):
            wb[t].column_dimensions[chr(65 + i)].width = width_list[i]
        for j in range(len(wb[t]['B'])):
            wb[t].row_dimensions[j + 1].height = 49.5
            for k in range(6):
                wb[t].cell(row=j + 1, column=1 + k).alignment = Alignment(
                    horizontal='center', vertical='center', wrap_text=True)
                wb[t].cell(row=j + 1, column=1 + k).font = fontStyle
            for k in range(3):
                wb[t].cell(row=j + 1, column=2 + k).alignment = Alignment(horizontal='left', wrap_text=True)
            for m in range(6):
                wb[t].cell(row=j + 1, column=1 + m).border = thin_border
    wb.save(file_path)


if __name__ == '__main__':
    main()
